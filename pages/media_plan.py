"""Media Plan page."""
import calendar
import io
import json
import os
from datetime import date, timedelta

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
import streamlit.components.v1 as _components
import toml

# ── Constants ─────────────────────────────────────────────────────────────────

MARKET_LABELS = {
    'AT': 'Austria',
    'BE': 'Belgium',
    'BG': 'Bulgaria',
    'HR': 'Croatia',
    'CY': 'Cyprus',
    'CZ': 'Czech Republic',
    'DK': 'Denmark',
    'EE': 'Estonia',
    'FI': 'Finland',
    'FR': 'France',
    'DE': 'Germany',
    'GR': 'Greece',
    'HU': 'Hungary',
    'IE': 'Ireland',
    'IT': 'Italy',
    'LV': 'Latvia',
    'LT': 'Lithuania',
    'LU': 'Luxembourg',
    'MT': 'Malta',
    'NL': 'Netherlands',
    'NO': 'Norway',
    'PL': 'Poland',
    'PT': 'Portugal',
    'RO': 'Romania',
    'SK': 'Slovakia',
    'SI': 'Slovenia',
    'ES': 'Spain',
    'SE': 'Sweden',
    'CH': 'Switzerland',
    'UK': 'United Kingdom',
}

# Default benchmark values — markets with known data have specific values,
# others fall back to regional averages (Western EU vs Eastern EU).
# CPM YouTube | CPM LinkedIn | View Rate | CTR YT | CTR LI | Frequency
# CPC Search  | CTR Search   | Click→Session | Conv Rate

def _default_bench(cpm_yt, cpm_li, view_rate, ctr_yt, ctr_li, freq,
                   cpc_s, ctr_s, c2s_yt=0.80, c2s_li=0.82, c2s_s=0.86,
                   cr=0.02, l2m=0.20, m2s=0.30):
    shared = {'conv_rate': cr, 'lead_to_mql': l2m, 'mql_to_sql': m2s}
    # Display CPM is ~30% of YouTube; CTR is much lower (~0.15%) due to banner format;
    # click intent is lower so click_to_session is also lower (~0.70).
    cpm_dis = round(cpm_yt * 0.30, 1)
    return {
        'YouTube':  {'cpm': cpm_yt, 'view_rate': view_rate, 'ctr': ctr_yt,
                     'frequency': freq, 'click_to_session': c2s_yt, **shared},
        'LinkedIn': {'cpm': cpm_li, 'ctr': ctr_li,
                     'frequency': freq, 'click_to_session': c2s_li, **shared},
        'Search':   {'cpc': cpc_s,  'ctr': ctr_s,
                     'click_to_session': c2s_s, **shared},
        'Display':  {'cpm': cpm_dis, 'ctr': 0.0015,
                     'frequency': 4.0, 'click_to_session': 0.70, **shared},
    }

BENCH = {
    'AT': _default_bench(11.0, 18.0, 0.31, 0.0035, 0.0038, 3.0, 2.30, 0.030),
    'BE': _default_bench(11.0, 17.0, 0.31, 0.0035, 0.0038, 3.0, 2.20, 0.030),
    'BG': _default_bench(3.5,  7.0,  0.30, 0.0028, 0.0032, 3.0, 0.70, 0.022),
    'HR': _default_bench(4.5,  9.0,  0.30, 0.0030, 0.0033, 3.0, 0.90, 0.023),
    'CY': _default_bench(7.0, 13.0,  0.30, 0.0030, 0.0035, 3.0, 1.50, 0.026),
    'CZ': _default_bench(5.0,  9.5,  0.30, 0.0030, 0.0033, 3.0, 1.00, 0.024),
    'DK': _default_bench(13.0, 21.0, 0.31, 0.0035, 0.0040, 3.0, 2.80, 0.032),
    'EE': _default_bench(4.5,  9.0,  0.30, 0.0030, 0.0033, 3.0, 0.90, 0.023),
    'FI': _default_bench(12.0, 19.0, 0.31, 0.0035, 0.0038, 3.0, 2.50, 0.030),
    'FR': _default_bench(10.0, 17.0, 0.31, 0.0033, 0.0038, 3.0, 2.00, 0.028),
    'DE': _default_bench(11.0, 18.0, 0.31, 0.0035, 0.0038, 3.0, 2.20, 0.030),
    'GR': _default_bench(5.0,  10.0, 0.30, 0.0028, 0.0032, 3.0, 1.00, 0.024),
    'HU': _default_bench(4.0,  8.0,  0.30, 0.0028, 0.0032, 3.0, 0.80, 0.022),
    'IE': _default_bench(12.0, 19.0, 0.31, 0.0035, 0.0040, 3.0, 2.50, 0.032),
    'IT': _default_bench(8.0,  14.0, 0.30, 0.0030, 0.0035, 3.0, 1.60, 0.026),
    'LV': _default_bench(4.5,  9.0,  0.30, 0.0030, 0.0033, 3.0, 0.90, 0.023),
    'LT': _default_bench(4.5,  9.0,  0.30, 0.0030, 0.0033, 3.0, 0.90, 0.023),
    'LU': _default_bench(12.0, 18.0, 0.31, 0.0035, 0.0038, 3.0, 2.40, 0.030),
    'MT': _default_bench(7.0,  13.0, 0.30, 0.0030, 0.0035, 3.0, 1.40, 0.025),
    'NL': _default_bench(10.0, 16.0, 0.31, 0.0035, 0.0040, 3.0, 2.00, 0.030),
    'NO': _default_bench(13.0, 22.0, 0.31, 0.0035, 0.0040, 3.0, 2.80, 0.032),
    'PL': _default_bench( 5.0,  9.0, 0.30, 0.0030, 0.0035, 3.0, 1.00, 0.025),
    'PT': _default_bench( 6.0, 11.0, 0.30, 0.0030, 0.0035, 3.0, 1.20, 0.025),
    'RO': _default_bench( 3.5,  7.0, 0.30, 0.0028, 0.0032, 3.0, 0.70, 0.022),
    'SK': _default_bench( 4.5,  9.0, 0.30, 0.0030, 0.0033, 3.0, 0.90, 0.023),
    'SI': _default_bench( 5.0, 10.0, 0.30, 0.0030, 0.0033, 3.0, 1.00, 0.024),
    'ES': _default_bench( 6.0, 12.0, 0.30, 0.0030, 0.0035, 3.0, 1.50, 0.025),
    'SE': _default_bench(12.0, 20.0, 0.31, 0.0035, 0.0040, 3.0, 2.60, 0.031),
    'CH': _default_bench(13.0, 21.0, 0.31, 0.0035, 0.0040, 3.0, 2.90, 0.032),
    'UK': _default_bench(12.0, 20.0, 0.31, 0.0035, 0.0040, 3.0, 2.50, 0.035),
}

CH_COLORS = {
    'YouTube':  ['#1F497D', '#437CA3', '#6B9FBF', '#9FC3D5', '#C5DCE8'],
    'LinkedIn': ['#1F6152', '#2E8A72', '#4DB896', '#7DCFB0', '#A8E4D0'],
    'Search':   ['#4285F4', '#5A95F5', '#74A5F6', '#8EB5F7', '#A8C5F8'],
    'Display':  ['#B45309', '#D97706', '#F59E0B', '#FCD34D', '#FEF3C7'],
}

ALL_GOALS = ['Awareness', 'Traffic', 'Conversion']

MARKET_GROUPS = {
    'DACH':        ['DE', 'AT', 'CH'],
    'Nordics':     ['DK', 'FI', 'NO', 'SE'],
    'BeNeLux':     ['BE', 'NL', 'LU'],
    'Southern EU': ['ES', 'IT', 'PT', 'GR'],
    'CEE':         ['PL', 'CZ', 'HU', 'RO', 'BG', 'SK', 'SI', 'HR'],
    'UK + IE':     ['UK', 'IE'],
    'All EU':      ['DE','AT','CH','DK','FI','NO','SE','BE','NL','LU','ES','IT','PT','GR',
                    'PL','CZ','HU','RO','BG','SK','SI','HR','FR','UK','IE','EE','LV','LT',
                    'CY','MT'],
}

PLAN_TEMPLATES = {
    'DACH Awareness Launch': {
        'markets': ['DE','AT','CH'], 'budget': 50000,
        'goals': {'Awareness': ['YouTube']},
    },
    'Pan-EU Lead Gen': {
        'markets': ['DE','FR','NL','BE','ES','IT','SE','PL'], 'budget': 120000,
        'goals': {'Traffic': ['LinkedIn','Search'], 'Conversion': ['Search']},
    },
    'Single Market Full Funnel': {
        'markets': ['DE'], 'budget': 80000,
        'goals': {'Awareness': ['YouTube'], 'Traffic': ['LinkedIn','Search'], 'Conversion': ['Search']},
    },
    'Nordics Brand Building': {
        'markets': ['DK','FI','NO','SE'], 'budget': 60000,
        'goals': {'Awareness': ['YouTube','LinkedIn']},
    },
    'BeNeLux Performance': {
        'markets': ['BE','NL','LU'], 'budget': 40000,
        'goals': {'Traffic': ['Search','LinkedIn'], 'Conversion': ['Search']},
    },
}

BENCH_PRESET_FACTORS = {
    'Conservative': {'cpm': 1.15, 'cpc': 1.15, 'ctr': 0.80, 'view_rate': 0.85, 'conv_rate': 0.75},
    'Average':      {'cpm': 1.00, 'cpc': 1.00, 'ctr': 1.00, 'view_rate': 1.00, 'conv_rate': 1.00},
    'Aggressive':   {'cpm': 0.88, 'cpc': 0.88, 'ctr': 1.25, 'view_rate': 1.15, 'conv_rate': 1.30},
}

BENCH_HELP = {
    'cpm':                  'Cost per 1,000 impressions. Western EU: €10–13. Eastern EU: €3.5–5.',
    'cpc':                  'Cost per click on Search. Ranges €0.70 (Eastern EU) to €2.90 (DACH/Nordics).',
    'ctr':                  'Click-through rate — % of ad impressions that result in a click.',
    'view_rate':            'YouTube: % of impressions resulting in a 30-second (or full-video) view.',
    'frequency':            'Average number of times one unique user sees your ad across the campaign.',
    'click_to_session':     '% of clicks that result in a tracked site session (accounts for pixel gaps and bounces).',
    'conv_rate':            '% of sessions that convert to a lead. Typical B2B range: 1–5%.',
    'lead_to_mql':          '% of raw leads that qualify as Marketing Qualified Leads. Typical B2B: 15–35%.',
    'mql_to_sql':           '% of MQLs accepted by sales as Sales Qualified Leads. Typical B2B: 20–40%.',
    'open_rate':            'Sponsored Message / Conversation Ad: % of sends that are opened.',
    'form_completion_rate': 'Lead Gen Form: % of ad clicks that complete and submit the LinkedIn form.',
}

DONUT_PALETTE = ['#2BB5A5','#1F497D','#4DB896','#437CA3','#5A3E7A','#E8A838','#8B3A3A','#6B7280',
                 '#229990','#2E8A72','#6B9FBF','#7DCFB0']

ADDITIVE = ['Budget', 'impressions', 'reach', 'views', 'clicks', 'sessions',
            'conversions', 'mql', 'sql',
            'sends', 'opens', 'cta_clicks', 'form_completions']

COL_FMT = {
    'Budget':           ('Spent (€)',        lambda x: f'€{x:,.0f}'),
    'impressions':      ('Impressions',      lambda x: f'{int(round(x)):,}'),
    'eff_cpm':          ('CPM (€)',          lambda x: f'€{x:.2f}'),
    'reach':            ('Reach',            lambda x: f'{int(round(x)):,}'),
    'views':            ('Views',            lambda x: f'{int(round(x)):,}'),
    'clicks':           ('Clicks',           lambda x: f'{int(round(x)):,}'),
    'cpc':              ('CPC (€)',          lambda x: f'€{x:.2f}'),
    'ctr':              ('CTR',              lambda x: f'{x*100:.2f}%'),
    'click_to_session': ('Click→Session %',  lambda x: f'{x*100:.0f}%'),
    'sessions':         ('Sessions',         lambda x: f'{int(round(x)):,}'),
    'conv_rate':        ('Session→Lead %',   lambda x: f'{x*100:.2f}%'),
    'conversions':      ('Leads',            lambda x: f'{int(round(x)):,}'),
    'cpa':              ('Cost per Lead (€)', lambda x: f'€{x:,.2f}'),
    'lead_to_mql':      ('Lead→MQL %',       lambda x: f'{x*100:.0f}%'),
    'mql':              ('MQL',              lambda x: f'{int(round(x)):,}'),
    'cost_per_mql':     ('CPMQL (€)',        lambda x: f'€{x:,.2f}'),
    'mql_to_sql':       ('MQL→SQL %',        lambda x: f'{x*100:.0f}%'),
    'sql':              ('SQL',              lambda x: f'{int(round(x)):,}'),
    'cost_per_sql':     ('CPSQL (€)',        lambda x: f'€{x:,.2f}'),
    'view_rate':            ('View Rate',              lambda x: f'{x*100:.1f}%'),
    'cpv':                  ('CPV (€)',                lambda x: f'€{x:.2f}'),
    'cvr':                  ('Click→Lead %',           lambda x: f'{x*100:.2f}%'),
    'click_to_session_raw': ('Click→Session',          lambda x: f'{x*100:.0f}%'),
    'sends':                ('Sends',                  lambda x: f'{int(round(x)):,}'),
    'open_rate':            ('Open Rate',              lambda x: f'{x*100:.1f}%'),
    'opens':                ('Opens',                  lambda x: f'{int(round(x)):,}'),
    'cta_clicks':           ('CTA Clicks',             lambda x: f'{int(round(x)):,}'),
    'form_completions':     ('Form Completions',       lambda x: f'{int(round(x)):,}'),
    'form_completion_rate': ('Form Completion %',      lambda x: f'{x*100:.1f}%'),
    'cost_per_send':        ('Cost per Send (€)',      lambda x: f'€{x:.4f}'),
    'cost_per_open':        ('Cost per Open (€)',      lambda x: f'€{x:.4f}'),
    'cta_ctr':              ('CTA CTR',                lambda x: f'{x*100:.2f}%'),
}

# Columns shown per (channel, goal), in the exact order the user specified.
_TRAFFIC_COLS    = ['Budget', 'impressions', 'eff_cpm', 'clicks', 'cpc',
                    'ctr', 'click_to_session', 'sessions']
_CONVERSION_COLS = ['Budget', 'impressions', 'eff_cpm', 'clicks', 'cpc',
                    'ctr', 'click_to_session', 'sessions',
                    'conv_rate', 'conversions', 'cpa',
                    'lead_to_mql', 'mql', 'cost_per_mql',
                    'mql_to_sql', 'sql', 'cost_per_sql']

PHASE_COLS = {
    ('YouTube',  'Awareness'):  ['Budget', 'impressions', 'reach', 'views', 'cpv', 'ctr', 'clicks', 'cpc'],
    ('YouTube',  'Traffic'):    _TRAFFIC_COLS,
    ('YouTube',  'Conversion'): _CONVERSION_COLS,
    ('Search',   'Awareness'):  ['Budget', 'impressions', 'ctr', 'clicks', 'cpc'],
    ('Search',   'Traffic'):    _TRAFFIC_COLS,
    ('Search',   'Conversion'): _CONVERSION_COLS,
    ('LinkedIn', 'Awareness'):  ['Budget', 'impressions', 'reach', 'ctr', 'clicks', 'cpc'],
    ('LinkedIn', 'Traffic'):    _TRAFFIC_COLS,
    ('LinkedIn', 'Conversion'): _CONVERSION_COLS,
    # LinkedIn Sponsored Message / Conversational Ad — no sessions, all within LinkedIn
    ('LinkedIn_SM', 'Awareness'):  ['Budget', 'sends', 'cost_per_send', 'opens', 'cost_per_open',
                                    'open_rate', 'cta_clicks', 'ctr'],
    ('LinkedIn_SM', 'Traffic'):    ['Budget', 'sends', 'cost_per_send', 'opens', 'cost_per_open',
                                    'open_rate', 'cta_clicks', 'ctr'],
    ('LinkedIn_SM', 'Conversion'): ['Budget', 'sends', 'cost_per_send', 'opens', 'cost_per_open',
                                    'open_rate', 'cta_clicks', 'ctr',
                                    'conversions', 'cpa',
                                    'lead_to_mql', 'mql', 'cost_per_mql',
                                    'mql_to_sql', 'sql', 'cost_per_sql'],
    # LinkedIn Conversation Ad — same metric structure as Sponsored Message
    ('LinkedIn_CA', 'Awareness'):  ['Budget', 'sends', 'cost_per_send', 'opens', 'cost_per_open',
                                    'open_rate', 'cta_clicks', 'ctr'],
    ('LinkedIn_CA', 'Traffic'):    ['Budget', 'sends', 'cost_per_send', 'opens', 'cost_per_open',
                                    'open_rate', 'cta_clicks', 'ctr'],
    ('LinkedIn_CA', 'Conversion'): ['Budget', 'sends', 'cost_per_send', 'opens', 'cost_per_open',
                                    'open_rate', 'cta_clicks', 'ctr',
                                    'conversions', 'cpa',
                                    'lead_to_mql', 'mql', 'cost_per_mql',
                                    'mql_to_sql', 'sql', 'cost_per_sql'],
    # LinkedIn Document Ad — document preview as creative, Lead Gen Form as conversion action
    ('LinkedIn_DA', 'Awareness'):  ['Budget', 'impressions', 'eff_cpm', 'clicks', 'ctr',
                                    'form_completions', 'form_completion_rate'],
    ('LinkedIn_DA', 'Traffic'):    ['Budget', 'impressions', 'eff_cpm', 'clicks', 'ctr',
                                    'form_completions', 'form_completion_rate'],
    ('LinkedIn_DA', 'Conversion'): ['Budget', 'impressions', 'eff_cpm', 'clicks', 'ctr',
                                    'form_completions', 'form_completion_rate',
                                    'lead_to_mql', 'mql', 'cost_per_mql',
                                    'mql_to_sql', 'sql', 'cost_per_sql'],
    # LinkedIn Lead Gen Form
    ('LinkedIn_LGF', 'Awareness'):  ['Budget', 'impressions', 'eff_cpm', 'clicks', 'ctr',
                                     'form_completions', 'form_completion_rate'],
    ('LinkedIn_LGF', 'Traffic'):    ['Budget', 'impressions', 'eff_cpm', 'clicks', 'ctr',
                                     'form_completions', 'form_completion_rate'],
    ('LinkedIn_LGF', 'Conversion'): ['Budget', 'impressions', 'eff_cpm', 'clicks', 'ctr',
                                     'form_completions', 'form_completion_rate',
                                     'lead_to_mql', 'mql', 'cost_per_mql',
                                     'mql_to_sql', 'sql', 'cost_per_sql'],
    # Display: banner/image — no views/cpv; reach via impressions/frequency
    ('Display',  'Awareness'):  ['Budget', 'impressions', 'reach', 'eff_cpm', 'ctr', 'clicks', 'cpc'],
    ('Display',  'Traffic'):    _TRAFFIC_COLS,
    ('Display',  'Conversion'): _CONVERSION_COLS,
}


# ── Helpers ───────────────────────────────────────────────────────────────────

def generate_periods(start, end, breakdown):
    periods, cur = [], start
    if breakdown == 'Daily':
        while cur <= end:
            periods.append({'label': cur.strftime('%b %d, %Y'), 'days': 1})
            cur += timedelta(days=1)
    elif breakdown == 'Weekly':
        while cur <= end:
            p_end = min(cur + timedelta(days=6), end)
            periods.append({'label': f"{cur.strftime('%b %d')} – {p_end.strftime('%b %d')}", 'days': (p_end - cur).days + 1})
            cur += timedelta(days=7)
    elif breakdown == 'Bi-Weekly':
        while cur <= end:
            p_end = min(cur + timedelta(days=13), end)
            periods.append({'label': f"{cur.strftime('%b %d')} – {p_end.strftime('%b %d')}", 'days': (p_end - cur).days + 1})
            cur += timedelta(days=14)
    elif breakdown == 'Monthly':
        while cur <= end:
            last = calendar.monthrange(cur.year, cur.month)[1]
            p_end = min(date(cur.year, cur.month, last), end)
            periods.append({'label': cur.strftime('%B %Y'), 'days': (p_end - cur).days + 1})
            cur = date(cur.year + 1, 1, 1) if cur.month == 12 else date(cur.year, cur.month + 1, 1)
    return periods


def calc_row(budget, bm, goal, channel, conv_rate):
    if budget <= 0:
        return {'Budget': budget}
    r = {'Budget': budget}

    if channel == 'Search':
        cpc = bm.get('cpc', 2.0)
        ctr = bm.get('ctr', 0.03)
        c2s = bm.get('click_to_session', 0.85)
        if cpc <= 0:
            return r
        clicks = budget / cpc
        impressions = clicks / ctr if ctr > 0 else 0
        r.update({'impressions': impressions, 'clicks': clicks, 'cpc': cpc, 'ctr': ctr})
        if goal in ('Traffic', 'Conversion'):
            r['sessions'] = clicks * c2s
            r['click_to_session'] = c2s
        if goal == 'Conversion':
            convs = r['sessions'] * conv_rate
            r['conversions'] = convs
            r['cpa']       = budget / convs if convs > 0 else 0
            r['cvr']       = convs / clicks if clicks > 0 else 0
            r['conv_rate'] = conv_rate
            l2m = bm.get('lead_to_mql', 0.20)
            m2s = bm.get('mql_to_sql',  0.30)
            mql = convs * l2m
            sql = mql * m2s
            r['lead_to_mql']  = l2m
            r['mql']          = mql
            r['cost_per_mql'] = budget / mql if mql > 0 else 0
            r['mql_to_sql']   = m2s
            r['sql']          = sql
            r['cost_per_sql'] = budget / sql if sql > 0 else 0

    elif channel == 'LinkedIn' and bm.get('_li_format') in (
            'Sponsored Message / Conversational Ad', 'Conversation Ad'):
        # Everything stays inside LinkedIn — no sessions metric
        cps       = bm.get('cpm', 0.50)    # repurposed field: cost per send
        open_rate = bm.get('open_rate', 0.30)
        cta_ctr   = bm.get('ctr', 0.10)    # CTA click rate out of opens
        if cps <= 0:
            return r
        sends      = budget / cps
        opens      = sends * open_rate
        cta_clicks = opens * cta_ctr
        r.update({
            'sends':         sends,
            'cost_per_send': cps,
            'opens':         opens,
            'cost_per_open': budget / opens if opens > 0 else 0,
            'open_rate':     open_rate,
            'cta_clicks':    cta_clicks,
            'ctr':           cta_ctr,
        })
        if goal == 'Conversion':
            # Leads come directly from CTA clicks — no session step
            c2l = bm.get('conv_rate', 0.05)  # CTA click → lead rate
            convs = cta_clicks * c2l
            r['conversions'] = convs
            r['conv_rate']   = c2l
            r['cpa']         = budget / convs if convs > 0 else 0
            l2m = bm.get('lead_to_mql', 0.20)
            m2s = bm.get('mql_to_sql',  0.30)
            mql = convs * l2m
            sql = mql * m2s
            r['lead_to_mql']  = l2m
            r['mql']          = mql
            r['cost_per_mql'] = budget / mql if mql > 0 else 0
            r['mql_to_sql']   = m2s
            r['sql']          = sql
            r['cost_per_sql'] = budget / sql if sql > 0 else 0

    elif channel == 'LinkedIn' and bm.get('_li_format') in ('Document Ad', 'Lead Gen Form'):
        cpm = bm.get('cpm', 10.0)
        ctr = bm.get('ctr', 0.005)
        fcr = bm.get('form_completion_rate', 0.10)
        if cpm <= 0:
            return r
        imp    = (budget / cpm) * 1000
        clicks = imp * ctr
        form_completions = clicks * fcr
        r.update({
            'impressions':        imp,
            'clicks':             clicks,
            'ctr':                ctr,
            'form_completions':   form_completions,
            'form_completion_rate': fcr,
        })
        if goal == 'Conversion':
            l2m = bm.get('lead_to_mql', 0.60)
            m2s = bm.get('mql_to_sql',  0.30)
            mql = form_completions * l2m
            sql = mql * m2s
            r['lead_to_mql']  = l2m
            r['mql']          = mql
            r['cost_per_mql'] = budget / mql if mql > 0 else 0
            r['mql_to_sql']   = m2s
            r['sql']          = sql
            r['cost_per_sql'] = budget / sql if sql > 0 else 0

    else:
        cpm = bm.get('cpm', 10.0)
        ctr = bm.get('ctr', 0.003)
        freq = bm.get('frequency', 3.0)
        c2s = bm.get('click_to_session', 0.80)
        if cpm <= 0:
            return r
        imp = (budget / cpm) * 1000
        reach = imp / freq if freq > 0 else 0
        clicks = imp * ctr
        r.update({
            'impressions': imp,
            'reach':       reach,
            'clicks':      clicks,
            'ctr':         ctr,
            'cpc':         budget / clicks if clicks > 0 else 0,
        })
        if goal == 'Awareness' and channel == 'YouTube':
            vr = bm.get('view_rate', 0.31)
            views = imp * vr
            r['views'] = views
            r['cpv']   = budget / views if views > 0 else 0
        if goal in ('Traffic', 'Conversion'):
            r['sessions']         = clicks * c2s
            r['click_to_session'] = c2s
        if goal == 'Conversion':
            convs = r['sessions'] * conv_rate
            r['conversions'] = convs
            r['cpa']         = budget / convs if convs > 0 else 0
            r['cvr']         = convs / clicks if clicks > 0 else 0
            r['conv_rate']   = conv_rate
            l2m = bm.get('lead_to_mql', 0.20)
            m2s = bm.get('mql_to_sql',  0.30)
            mql = convs * l2m
            sql = mql * m2s
            r['lead_to_mql']  = l2m
            r['mql']          = mql
            r['cost_per_mql'] = budget / mql if mql > 0 else 0
            r['mql_to_sql']   = m2s
            r['sql']          = sql
            r['cost_per_sql'] = budget / sql if sql > 0 else 0

    # Effective CPM — works for impression-based channels
    if r.get('impressions', 0) > 0:
        r['eff_cpm'] = r['Budget'] / r['impressions'] * 1000

    return r


def build_table(periods, total_budget, bm, goal, channel, conv_rate):
    total_days = sum(p['days'] for p in periods) or 1
    rows = []
    for p in periods:
        bud = total_budget * p['days'] / total_days
        m = calc_row(bud, bm, goal, channel, conv_rate)
        rows.append({'Period': p['label'], 'Days': p['days'], **m})
    df = pd.DataFrame(rows)
    total_m = calc_row(total_budget, bm, goal, channel, conv_rate)
    total_df = pd.DataFrame([{'Period': 'TOTAL', 'Days': total_days, **total_m}])
    return pd.concat([df, total_df], ignore_index=True)


def fmt_df(df, ch=None, goal=None, li_format=None):
    out = df[['Period', 'Days']].copy()
    ch_key = ch
    if ch == 'LinkedIn' and li_format:
        if li_format == 'Sponsored Message / Conversational Ad':
            ch_key = 'LinkedIn_SM'
        elif li_format == 'Lead Gen Form':
            ch_key = 'LinkedIn_LGF'
    col_order = PHASE_COLS.get((ch_key, goal)) if ch_key and goal else None
    keys = col_order if col_order else list(COL_FMT.keys())
    for raw in keys:
        if raw in COL_FMT and raw in df.columns:
            label, fn = COL_FMT[raw]
            out[label] = df[raw].apply(fn)
    return out


def make_funnel(df, goal, channel, title, li_format=None):
    t = df[df['Period'] == 'TOTAL'].iloc[0]
    colors = CH_COLORS.get(channel, CH_COLORS['YouTube'])

    if channel == 'Search':
        if goal == 'Awareness':
            stages = [('Impressions', 'impressions'), ('Clicks', 'clicks')]
        elif goal == 'Traffic':
            stages = [('Impressions', 'impressions'), ('Clicks', 'clicks'), ('Sessions', 'sessions')]
        else:
            stages = [('Impressions', 'impressions'), ('Clicks', 'clicks'), ('Sessions', 'sessions'), ('Conversions', 'conversions')]
    elif channel == 'YouTube':
        if goal == 'Awareness':
            stages = [('Impressions', 'impressions'), ('Reach', 'reach'), ('Views', 'views'), ('Clicks', 'clicks')]
        elif goal == 'Traffic':
            stages = [('Impressions', 'impressions'), ('Clicks', 'clicks'), ('Sessions', 'sessions')]
        else:
            stages = [('Impressions', 'impressions'), ('Clicks', 'clicks'), ('Sessions', 'sessions'), ('Conversions', 'conversions')]
    elif channel == 'LinkedIn' and li_format in (
            'Sponsored Message / Conversational Ad', 'Conversation Ad'):
        if goal == 'Conversion':
            stages = [('Sends', 'sends'), ('Opens', 'opens'), ('CTA Clicks', 'cta_clicks'),
                      ('Leads', 'conversions'), ('MQLs', 'mql'), ('SQLs', 'sql')]
        else:
            stages = [('Sends', 'sends'), ('Opens', 'opens'), ('CTA Clicks', 'cta_clicks')]
    elif channel == 'LinkedIn' and li_format == 'Document Ad':
        if goal == 'Conversion':
            stages = [('Impressions', 'impressions'), ('Clicks', 'clicks'),
                      ('Form Completions', 'form_completions'), ('MQLs', 'mql'), ('SQLs', 'sql')]
        else:
            stages = [('Impressions', 'impressions'), ('Clicks', 'clicks'),
                      ('Form Completions', 'form_completions')]
    elif channel == 'LinkedIn' and li_format == 'Lead Gen Form':
        if goal == 'Conversion':
            stages = [('Impressions', 'impressions'), ('Clicks', 'clicks'),
                      ('Form Completions', 'form_completions'), ('MQLs', 'mql'), ('SQLs', 'sql')]
        else:
            stages = [('Impressions', 'impressions'), ('Clicks', 'clicks'), ('Form Completions', 'form_completions')]
    else:  # LinkedIn Standard / Display
        if goal == 'Awareness':
            stages = [('Impressions', 'impressions'), ('Reach', 'reach'), ('Clicks', 'clicks')]
        elif goal == 'Traffic':
            stages = [('Impressions', 'impressions'), ('Clicks', 'clicks'), ('Sessions', 'sessions')]
        else:
            stages = [('Impressions', 'impressions'), ('Clicks', 'clicks'), ('Sessions', 'sessions'), ('Conversions', 'conversions')]

    labels = [s[0] for s in stages]
    values = [t.get(s[1], 0) for s in stages]

    fig = go.Figure(go.Funnel(
        y=labels,
        x=values,
        textinfo='value+percent initial',
        marker={'color': colors[:len(stages)]},
        textfont={'size': 10},
        connector={'line': {'color': 'rgba(0,0,0,0.1)', 'width': 1}},
    ))
    fig.update_layout(
        title={'text': title, 'font': {'size': 12, 'color': '#1F497D'}, 'x': 0.5, 'xanchor': 'center'},
        margin={'t': 40, 'b': 10, 'l': 10, 'r': 10},
        height=265,
        paper_bgcolor='rgba(0,0,0,0)',
    )
    return fig


def _duplicate_scenario(sid):
    import re as _re
    new_sid = len(st.session_state.scenario_names)
    st.session_state.scenario_names.append(st.session_state.scenario_names[sid] + ' (copy)')
    skip_prefixes = ('remove_', 'dup_', 'rename_', 'btn_', 'dl_', 'funnel_', 'grand_',
                     'preset_', 'eq_', 'cpm_eff_', 'grp_', 'tpl_', 'pacing_', 'bar_',
                     'save_tpl_', 'del_tpl_', 'pin_', 'mkt_tog_',
                     'bud_mkt_', '_last_total_', 'li_fmt_prev_')
    pattern = _re.compile(rf'^(.+)_{sid}$')
    for k, v in list(st.session_state.items()):
        m = pattern.match(str(k))
        if m and not any(str(k).startswith(p) for p in skip_prefixes):
            st.session_state[f'{m.group(1)}_{new_sid}'] = v
    # Don't call st.rerun() here — that would interrupt the mid-render tab loop,
    # causing Streamlit to orphan widget state for tabs not yet rendered. Instead
    # we set a flag and rerun AFTER the full tab loop finishes (see bottom of file).
    st.session_state['_pending_dup_rerun'] = True


def _apply_template_data(sid, tpl):
    """Apply a template dict to a scenario slot."""
    n = max(len(tpl['markets']), 1)
    st.session_state[f'selected_markets_{sid}'] = tpl['markets']
    st.session_state[f'total_budget_{sid}']     = tpl['budget']
    default_pct = round(100.0 / n, 1)
    for mkt in tpl['markets']:
        st.session_state[f'pct_{mkt}_{sid}'] = default_pct
    for goal in ALL_GOALS:
        goal_on = goal in tpl['goals']
        st.session_state[f'sb_goal_{goal}_{sid}'] = goal_on
        for ch, key in [('YouTube', 'sb_yt'), ('Search', 'sb_s'), ('LinkedIn', 'sb_li'), ('Display', 'sb_dis')]:
            st.session_state[f'{key}_{goal}_{sid}'] = goal_on and ch in tpl['goals'].get(goal, [])
    st.session_state['_pending_dup_rerun'] = True


def _apply_template(sid, tpl_name):
    _apply_template_data(sid, PLAN_TEMPLATES[tpl_name])


def _sync_pct_to_bud(mkt, sid):
    """on_change for % input — keeps the € input in sync."""
    total = st.session_state.get(f'total_budget_{sid}', 0)
    pct   = st.session_state.get(f'pct_{mkt}_{sid}', 0)
    st.session_state[f'bud_mkt_{mkt}_{sid}'] = round(total * pct / 100)


def _sync_bud_to_pct(mkt, sid):
    """on_change for € input — keeps the % input in sync."""
    total = st.session_state.get(f'total_budget_{sid}', 0)
    bud   = st.session_state.get(f'bud_mkt_{mkt}_{sid}', 0)
    st.session_state[f'pct_{mkt}_{sid}'] = round(bud / total * 100, 1) if total > 0 else 0.0


def _current_as_template(sid):
    """Snapshot the current scenario config as a saveable template dict."""
    ss = st.session_state
    goals = {}
    for goal in ALL_GOALS:
        if ss.get(f'sb_goal_{goal}_{sid}', False):
            chs = []
            if ss.get(f'sb_yt_{goal}_{sid}',  False): chs.append('YouTube')
            if ss.get(f'sb_s_{goal}_{sid}',   False): chs.append('Search')
            if ss.get(f'sb_li_{goal}_{sid}',  False): chs.append('LinkedIn')
            if ss.get(f'sb_dis_{goal}_{sid}', False): chs.append('Display')
            if chs:
                goals[goal] = chs
    return {
        'markets': ss.get(f'selected_markets_{sid}', []),
        'budget':  ss.get(f'total_budget_{sid}', 0),
        'goals':   goals,
    }


def _apply_bench_preset(ch, mkt, goal, sid, preset_name, li_fmt=None):
    b = BENCH[mkt][ch]
    f = BENCH_PRESET_FACTORS[preset_name]
    keys = []
    if ch == 'Search':
        keys = [
            (f'cpc_{mkt}_{ch}_{goal}_{sid}',              round(b['cpc'] * f['cpc'], 2)),
            (f'ctr_{mkt}_{ch}_{goal}_{sid}',              round(b['ctr'] * 100 * f['ctr'], 2)),
            (f'click_to_session_{mkt}_{ch}_{goal}_{sid}', round(b.get('click_to_session', 0.85) * 100, 0)),
            (f'conv_rate_{mkt}_{ch}_{goal}_{sid}',        round(b.get('conv_rate', 0.03) * 100 * f['conv_rate'], 1)),
        ]
    elif ch == 'YouTube':
        keys = [
            (f'cpm_{mkt}_{ch}_{goal}_{sid}',              round(b['cpm'] * f['cpm'], 2)),
            (f'view_rate_{mkt}_{ch}_{goal}_{sid}',        round(b.get('view_rate', 0.31) * 100 * f['view_rate'], 1)),
            (f'ctr_{mkt}_{ch}_{goal}_{sid}',              round(b['ctr'] * 100 * f['ctr'], 2)),
            (f'click_to_session_{mkt}_{ch}_{goal}_{sid}', round(b.get('click_to_session', 0.80) * 100, 0)),
            (f'conv_rate_{mkt}_{ch}_{goal}_{sid}',        round(b.get('conv_rate', 0.02) * 100 * f['conv_rate'], 1)),
        ]
    elif ch == 'LinkedIn' and li_fmt == 'Sponsored Message / Conversational Ad':
        keys = [
            (f'cpm_{mkt}_{ch}_{goal}_{sid}',       round(b['cpm'] * f['cpm'] * 0.05, 3)),
            (f'open_rate_{mkt}_{ch}_{goal}_{sid}', round(30.0 * f.get('ctr', 1.0), 1)),
            (f'ctr_{mkt}_{ch}_{goal}_{sid}',       round(b['ctr'] * 100 * f['ctr'], 2)),
            (f'conv_rate_{mkt}_{ch}_{goal}_{sid}', round(5.0 * f['conv_rate'], 1)),
        ]
    elif ch == 'LinkedIn' and li_fmt == 'Conversation Ad':
        keys = [
            (f'cpm_{mkt}_{ch}_{goal}_{sid}',       round(b['cpm'] * f['cpm'] * 0.04, 3)),  # slightly cheaper than SM
            (f'open_rate_{mkt}_{ch}_{goal}_{sid}', round(45.0 * f.get('ctr', 1.0), 1)),
            (f'ctr_{mkt}_{ch}_{goal}_{sid}',       round(b['ctr'] * 100 * f['ctr'], 2)),
            (f'conv_rate_{mkt}_{ch}_{goal}_{sid}', round(6.0 * f['conv_rate'], 1)),
        ]
    elif ch == 'LinkedIn' and li_fmt == 'Document Ad':
        keys = [
            (f'cpm_{mkt}_{ch}_{goal}_{sid}',                  round(b['cpm'] * f['cpm'], 2)),
            (f'ctr_{mkt}_{ch}_{goal}_{sid}',                  round(b['ctr'] * 100 * f['ctr'], 2)),
            (f'form_completion_rate_{mkt}_{ch}_{goal}_{sid}', round(8.0 * f.get('conv_rate', 1.0), 1)),
            (f'lead_to_mql_{mkt}_{ch}_{goal}_{sid}',          50.0),
            (f'mql_to_sql_{mkt}_{ch}_{goal}_{sid}',           round(b.get('mql_to_sql', 0.30) * 100, 0)),
        ]
    elif ch == 'LinkedIn' and li_fmt == 'Lead Gen Form':
        keys = [
            (f'cpm_{mkt}_{ch}_{goal}_{sid}',                   round(b['cpm'] * f['cpm'], 2)),
            (f'ctr_{mkt}_{ch}_{goal}_{sid}',                   round(b['ctr'] * 100 * f['ctr'], 2)),
            (f'form_completion_rate_{mkt}_{ch}_{goal}_{sid}',  round(10.0 * f.get('conv_rate', 1.0), 1)),
            (f'lead_to_mql_{mkt}_{ch}_{goal}_{sid}',           60.0),
            (f'mql_to_sql_{mkt}_{ch}_{goal}_{sid}',            round(b.get('mql_to_sql', 0.30) * 100, 0)),
        ]
    else:  # LinkedIn Standard / Display
        keys = [
            (f'cpm_{mkt}_{ch}_{goal}_{sid}',              round(b['cpm'] * f['cpm'], 2)),
            (f'ctr_{mkt}_{ch}_{goal}_{sid}',              round(b['ctr'] * 100 * f['ctr'], 2)),
            (f'click_to_session_{mkt}_{ch}_{goal}_{sid}', round(b.get('click_to_session', 0.70 if ch == 'Display' else 0.82) * 100, 0)),
            (f'conv_rate_{mkt}_{ch}_{goal}_{sid}',        round(b.get('conv_rate', 0.02) * 100 * f['conv_rate'], 1)),
        ]
    for k, v in keys:
        st.session_state[k] = v


_BENCH_FIELDS = {
    ('YouTube',       'Awareness'):  ['cpm', 'view_rate', 'ctr', 'frequency'],
    ('YouTube',       'Traffic'):    ['cpm', 'ctr', 'click_to_session'],
    ('YouTube',       'Conversion'): ['cpm', 'ctr', 'click_to_session', 'conv_rate', 'lead_to_mql', 'mql_to_sql'],
    ('LinkedIn',      'Awareness'):  ['cpm', 'ctr', 'frequency'],
    ('LinkedIn',      'Traffic'):    ['cpm', 'ctr', 'click_to_session'],
    ('LinkedIn',      'Conversion'): ['cpm', 'ctr', 'click_to_session', 'conv_rate', 'lead_to_mql', 'mql_to_sql'],
    ('LinkedIn_SM',   'Awareness'):  ['cpm', 'open_rate', 'ctr'],
    ('LinkedIn_SM',   'Traffic'):    ['cpm', 'open_rate', 'ctr'],
    ('LinkedIn_SM',   'Conversion'): ['cpm', 'open_rate', 'ctr', 'conv_rate', 'lead_to_mql', 'mql_to_sql'],
    ('LinkedIn_CA',   'Awareness'):  ['cpm', 'open_rate', 'ctr'],
    ('LinkedIn_CA',   'Traffic'):    ['cpm', 'open_rate', 'ctr'],
    ('LinkedIn_CA',   'Conversion'): ['cpm', 'open_rate', 'ctr', 'conv_rate', 'lead_to_mql', 'mql_to_sql'],
    ('LinkedIn_DA',   'Awareness'):  ['cpm', 'ctr', 'form_completion_rate'],
    ('LinkedIn_DA',   'Traffic'):    ['cpm', 'ctr', 'form_completion_rate'],
    ('LinkedIn_DA',   'Conversion'): ['cpm', 'ctr', 'form_completion_rate', 'lead_to_mql', 'mql_to_sql'],
    ('LinkedIn_LGF',  'Awareness'):  ['cpm', 'ctr', 'form_completion_rate'],
    ('LinkedIn_LGF',  'Traffic'):    ['cpm', 'ctr', 'form_completion_rate'],
    ('LinkedIn_LGF',  'Conversion'): ['cpm', 'ctr', 'form_completion_rate', 'lead_to_mql', 'mql_to_sql'],
    ('Search',        'Awareness'):  ['cpc', 'ctr'],
    ('Search',        'Traffic'):    ['cpc', 'ctr', 'click_to_session'],
    ('Search',        'Conversion'): ['cpc', 'ctr', 'click_to_session', 'conv_rate', 'lead_to_mql', 'mql_to_sql'],
    ('Display',       'Awareness'):  ['cpm', 'ctr', 'frequency'],
    ('Display',       'Traffic'):    ['cpm', 'ctr', 'click_to_session'],
    ('Display',       'Conversion'): ['cpm', 'ctr', 'click_to_session', 'conv_rate', 'lead_to_mql', 'mql_to_sql'],
}

_BENCH_FIELD_DESC = {
    'cpm':                  'Cost per thousand impressions in EUR (e.g. 12.5). For Sponsored Message: cost per send in EUR (e.g. 0.50)',
    'cpc':                  'Cost per click in EUR (e.g. 2.5)',
    'ctr':                  'Click-through rate as decimal proportion — NOT percentage (e.g. 0.003 for 0.3%)',
    'view_rate':            'Video view/completion rate as decimal proportion (e.g. 0.30 for 30%)',
    'frequency':            'Average ad exposures per unique user over the campaign flight (e.g. 3.0)',
    'click_to_session':     'Post-click session rate as decimal proportion (e.g. 0.80 for 80%)',
    'conv_rate':            'Session-to-lead conversion rate as decimal proportion (e.g. 0.02 for 2%)',
    'lead_to_mql':          'Lead-to-MQL rate as decimal proportion (e.g. 0.20 for 20%)',
    'mql_to_sql':           'MQL-to-SQL rate as decimal proportion (e.g. 0.30 for 30%)',
    'open_rate':            'Sponsored/Conversation Message open rate as decimal proportion (e.g. 0.35 for 35%)',
    'form_completion_rate': 'Lead Gen Form / Document Ad: % of clicks that complete the form, as decimal (e.g. 0.08 for 8%)',
}

_PRESET_DESC = {
    'Conservative': 'pessimistic but realistic — higher costs, lower engagement — safe for budget planning',
    'Average':      'typical industry benchmark — most likely outcome based on current market conditions',
    'Aggressive':   'optimistic stretch target — lower costs, higher engagement — best-case realistic scenario',
}

_BENCH_IS_PCT = {'ctr', 'view_rate', 'click_to_session', 'conv_rate', 'lead_to_mql', 'mql_to_sql',
                 'open_rate', 'form_completion_rate'}


def get_api_key():
    try:
        key = st.secrets['anthropic'].get('api_key') or st.secrets['anthropic'].get('ANTHROPIC_API_KEY', '')
        if key:
            return key
    except Exception:
        pass
    local = os.path.join(os.path.dirname(__file__), '.streamlit', 'secrets.toml')
    try:
        s = toml.load(local)
        return s['anthropic'].get('api_key') or s['anthropic'].get('ANTHROPIC_API_KEY', '')
    except Exception:
        return ''


def _apply_bench_preset_ai(ch, mkt, goal, sid, preset_name, audience, industry, api_key, li_fmt=None):
    """Call Claude to get industry/audience-specific benchmarks. Returns True on success."""
    import anthropic as _anthropic

    # Choose the right field list based on LinkedIn format
    if ch == 'LinkedIn' and li_fmt == 'Sponsored Message / Conversational Ad':
        bench_key = ('LinkedIn_SM', goal)
    elif ch == 'LinkedIn' and li_fmt == 'Conversation Ad':
        bench_key = ('LinkedIn_CA', goal)
    elif ch == 'LinkedIn' and li_fmt == 'Document Ad':
        bench_key = ('LinkedIn_DA', goal)
    elif ch == 'LinkedIn' and li_fmt == 'Lead Gen Form':
        bench_key = ('LinkedIn_LGF', goal)
    else:
        bench_key = (ch, goal)
    fields = _BENCH_FIELDS.get(bench_key, ['cpm', 'ctr'])
    properties = {f: {'type': 'number', 'description': _BENCH_FIELD_DESC[f]} for f in fields}

    ch_label = ch
    if ch == 'LinkedIn' and li_fmt:
        ch_label = f'LinkedIn ({li_fmt})'

    tool_def = {
        'name': 'set_benchmarks',
        'description': 'Return calibrated benchmark values for the specified channel/goal/audience/industry combination.',
        'input_schema': {'type': 'object', 'properties': properties, 'required': fields},
    }

    prompt = (
        f'You are a digital advertising benchmark expert specialising in European paid media.\n'
        f'Provide **{preset_name}** benchmark values for:\n'
        f'- Channel: {ch_label}\n'
        f'- Phase / Goal: {goal}\n'
        f'- Audience: {audience}\n'
        f'- Industry: {industry}\n'
        f'- Market: {MARKET_LABELS[mkt]}\n\n'
        f'{preset_name} means: {_PRESET_DESC[preset_name]}.\n\n'
        f'Base your answer on 2025-2026 European digital advertising benchmarks '
        f'specific to this audience type and industry. '
        f'Adjust for the market — Western EU (DE, UK, FR, NL) costs more than Eastern EU (PL, RO, BG). '
        f'Use the set_benchmarks tool to return the values. '
        f'All rate fields must be decimal proportions, NOT percentages.'
    )

    try:
        client = _anthropic.Anthropic(api_key=api_key)
        resp = client.messages.create(
            model='claude-haiku-4-5-20251001',
            max_tokens=256,
            tools=[tool_def],
            tool_choice={'type': 'tool', 'name': 'set_benchmarks'},
            messages=[{'role': 'user', 'content': prompt}],
        )
        block = next((b for b in resp.content if b.type == 'tool_use'), None)
        if not block:
            return False
        for field, raw_val in block.input.items():
            if field not in fields:
                continue
            v = float(raw_val)
            sk = f'{field}_{mkt}_{ch}_{goal}_{sid}'
            if field in _BENCH_IS_PCT:
                st.session_state[sk] = round(v * 100, 3)
            elif field == 'frequency':
                st.session_state[sk] = round(v, 1)
            else:
                st.session_state[sk] = round(v, 2)
        return True
    except Exception:
        return False


def _scenario_status(sid):
    ss = st.session_state
    has_goals   = any(ss.get(f'sb_goal_{g}_{sid}', False) for g in ALL_GOALS)
    has_markets = bool(ss.get(f'selected_markets_{sid}', []))
    has_budget  = ss.get(f'total_budget_{sid}', 0) > 0
    mkts        = ss.get(f'selected_markets_{sid}', [])
    pct_sum     = sum(ss.get(f'pct_{m}_{sid}', 0) for m in mkts)
    split_ok    = has_markets and abs(pct_sum - 100) <= 0.5
    steps = [
        ('Goals',   has_goals),
        ('Markets', has_markets),
        ('Budget',  has_budget),
        ('Split',   split_ok),
    ]
    parts = []
    for label, ok in steps:
        color = '#2BB5A5' if ok else '#E8A838'
        icon  = '✓' if ok else '○'
        parts.append(
            f'<span style="background:{color};color:white;border-radius:4px;'
            f'padding:2px 8px;font-size:0.72rem;font-weight:600;margin-right:4px">'
            f'{icon} {label}</span>'
        )
    st.markdown('<div style="margin-bottom:8px">' + ''.join(parts) + '</div>', unsafe_allow_html=True)


def _market_donut(market_pcts):
    labels = [f'{MARKET_LABELS[m]}' for m in market_pcts]
    values = list(market_pcts.values())
    colors = DONUT_PALETTE[:len(labels)]
    fig = go.Figure(go.Pie(
        labels=labels, values=values, hole=0.55,
        marker={'colors': colors, 'line': {'color': 'white', 'width': 2}},
        textinfo='percent', textfont={'size': 10},
        hovertemplate='%{label}: %{value:.1f}%<extra></extra>',
    ))
    fig.update_layout(
        margin={'t': 10, 'b': 10, 'l': 10, 'r': 10},
        height=200,
        paper_bgcolor='rgba(0,0,0,0)',
        showlegend=True,
        legend={'font': {'size': 12}, 'orientation': 'v', 'x': 1.02},
    )
    return fig


def _pacing_chart(periods, budget, sid):
    if not periods or budget <= 0:
        return
    total_days = sum(p['days'] for p in periods) or 1
    labels = [p['label'] for p in periods]
    values = [round(budget * p['days'] / total_days, 0) for p in periods]
    fig = go.Figure(go.Bar(
        x=labels, y=values,
        marker_color='#2BB5A5',
        text=[f'€{v:,.0f}' for v in values],
        textposition='outside',
        textfont={'size': 12, 'color': '#1A1A1A'},
    ))
    fig.update_layout(
        title={'text': 'Budget Pacing', 'font': {'size': 13, 'color': '#1A1A1A'}, 'x': 0.5, 'xanchor': 'center'},
        margin={'t': 40, 'b': 30, 'l': 10, 'r': 10},
        height=240,
        paper_bgcolor='rgba(0,0,0,0)',
        yaxis={'showticklabels': False, 'showgrid': False, 'zeroline': False},
        xaxis={'tickangle': -30 if len(periods) > 6 else 0, 'tickfont': {'size': 11}},
    )
    st.plotly_chart(fig, use_container_width=True, config={'displayModeBar': False}, key=f'pacing_{sid}')


_LI_ALL_BENCH_KEYS = [
    'cpm', 'ctr', 'frequency', 'click_to_session', 'conv_rate',
    'lead_to_mql', 'mql_to_sql', 'open_rate', 'form_completion_rate',
]

def _clear_li_bench_keys(mkt, goal, sid):
    for field in _LI_ALL_BENCH_KEYS:
        k = f'{field}_{mkt}_LinkedIn_{goal}_{sid}'
        if k in st.session_state:
            del st.session_state[k]


_LI_FORMATS = [
    'Static (Single Image)',
    'Video',
    'Carousel',
    'Sponsored Message / Conversational Ad',
    'Conversation Ad',
    'Document Ad',
    'Lead Gen Form',
]


def benchmark_inputs(ch, mkt, goal, sid=0):
    """Render benchmark inputs for one channel/market/goal. Returns a benchmark dict."""
    b = BENCH[mkt][ch]

    # LinkedIn: format dropdown comes first so preset buttons use the right fields
    li_fmt = None
    if ch == 'LinkedIn':
        li_fmt_key      = f'li_fmt_{mkt}_{goal}_{sid}'
        li_fmt_prev_key = f'li_fmt_prev_{mkt}_{goal}_{sid}'
        st.selectbox('LinkedIn Format', _LI_FORMATS, key=li_fmt_key)
        li_fmt     = st.session_state[li_fmt_key]
        li_fmt_prev = st.session_state.get(li_fmt_prev_key)
        if li_fmt_prev is not None and li_fmt_prev != li_fmt:
            _clear_li_bench_keys(mkt, goal, sid)
        st.session_state[li_fmt_prev_key] = li_fmt

    # Preset buttons — AI-powered when API key is available, generic fallback otherwise
    pc = st.columns([1, 1, 1, 4])
    _api_key_for_preset = get_api_key()
    _audience_for_preset  = st.session_state.get('audience_type', 'B2B')
    _industry_for_preset  = st.session_state.get('industry', 'Logistics, Supply Chain & Transportation')
    for i, pname in enumerate(['Conservative', 'Average', 'Aggressive']):
        if pc[i].button(pname, key=f'preset_{pname}_{ch}_{mkt}_{goal}_{sid}', use_container_width=True):
            if _api_key_for_preset:
                with st.spinner(f'{pname} · {_audience_for_preset} · {_industry_for_preset} · {ch} {goal} · {MARKET_LABELS[mkt]}…'):
                    ok = _apply_bench_preset_ai(
                        ch, mkt, goal, sid, pname,
                        _audience_for_preset, _industry_for_preset, _api_key_for_preset,
                        li_fmt=li_fmt,
                    )
                if not ok:
                    _apply_bench_preset(ch, mkt, goal, sid, pname, li_fmt=li_fmt)
            else:
                _apply_bench_preset(ch, mkt, goal, sid, pname, li_fmt=li_fmt)
            st.session_state['_pending_dup_rerun'] = True

    fields = []
    if ch == 'Search':
        fields = [
            ('cpc', 'CPC (€)',   b['cpc'],              0.10, '%.2f', False),
            ('ctr', 'CTR %',     b['ctr'] * 100,        0.05, '%.2f', True),
        ]
        if goal in ('Traffic', 'Conversion'):
            fields.append(('click_to_session', 'Click→Session %',
                           b.get('click_to_session', 0.85) * 100, 1.0, '%.0f', True))
        if goal == 'Conversion':
            fields += [
                ('conv_rate',   'Session→Lead %', b.get('conv_rate',   0.03) * 100, 0.1, '%.1f', True),
                ('lead_to_mql', 'Lead→MQL %',     b.get('lead_to_mql', 0.20) * 100, 1.0, '%.0f', True),
                ('mql_to_sql',  'MQL→SQL %',      b.get('mql_to_sql',  0.30) * 100, 1.0, '%.0f', True),
            ]

    elif ch == 'YouTube':
        fields = [('cpm', 'CPM (€)', b['cpm'], 0.5, '%.2f', False)]
        if goal == 'Awareness':
            fields.append(('view_rate', 'View Rate %',
                           b.get('view_rate', 0.31) * 100, 0.5, '%.1f', True))
        fields += [
            ('ctr',       'CTR %',     b['ctr'] * 100,          0.01, '%.2f', True),
            ('frequency', 'Frequency', b.get('frequency', 3.0), 0.5,  '%.1f', False),
        ]
        if goal in ('Traffic', 'Conversion'):
            fields.append(('click_to_session', 'Click→Session %',
                           b.get('click_to_session', 0.80) * 100, 1.0, '%.0f', True))
        if goal == 'Conversion':
            fields += [
                ('conv_rate',   'Session→Lead %', b.get('conv_rate',   0.02) * 100, 0.1, '%.1f', True),
                ('lead_to_mql', 'Lead→MQL %',     b.get('lead_to_mql', 0.20) * 100, 1.0, '%.0f', True),
                ('mql_to_sql',  'MQL→SQL %',      b.get('mql_to_sql',  0.30) * 100, 1.0, '%.0f', True),
            ]

    elif ch == 'LinkedIn':
        if li_fmt == 'Sponsored Message / Conversational Ad':
            fields = [
                ('cpm',       'Cost per Send (€)', b.get('cpm', 0.50), 0.01, '%.3f', False),
                ('open_rate', 'Open Rate %',        30.0,               0.5,  '%.1f', True),
                ('ctr',       'CTA Click Rate %',   b['ctr'] * 100,     0.1,  '%.2f', True),
            ]
            if goal == 'Conversion':
                fields += [
                    ('conv_rate',   'CTA Click→Lead %', 5.0,                              0.1, '%.1f', True),
                    ('lead_to_mql', 'Lead→MQL %',       b.get('lead_to_mql', 0.20) * 100, 1.0, '%.0f', True),
                    ('mql_to_sql',  'MQL→SQL %',        b.get('mql_to_sql',  0.30) * 100, 1.0, '%.0f', True),
                ]
        elif li_fmt == 'Conversation Ad':
            fields = [
                ('cpm',       'Cost per Send (€)', b.get('cpm', 0.40), 0.01, '%.3f', False),
                ('open_rate', 'Open Rate %',        45.0,               0.5,  '%.1f', True),
                ('ctr',       'CTA Click Rate %',   b['ctr'] * 100,     0.1,  '%.2f', True),
            ]
            if goal == 'Conversion':
                fields += [
                    ('conv_rate',   'CTA Click→Lead %', 6.0,                              0.1, '%.1f', True),
                    ('lead_to_mql', 'Lead→MQL %',       b.get('lead_to_mql', 0.20) * 100, 1.0, '%.0f', True),
                    ('mql_to_sql',  'MQL→SQL %',        b.get('mql_to_sql',  0.30) * 100, 1.0, '%.0f', True),
                ]
        elif li_fmt == 'Document Ad':
            # Document preview as creative, Lead Gen Form as the conversion action
            fields = [
                ('cpm',                  'CPM (€)',                b['cpm'],       0.1,  '%.2f', False),
                ('ctr',                  'CTR %',                  b['ctr'] * 100, 0.01, '%.3f', True),
                ('form_completion_rate', 'Form Completion Rate %', 8.0,            0.5,  '%.1f', True),
                ('lead_to_mql',          'Lead→MQL %',             50.0,           1.0,  '%.0f', True),
                ('mql_to_sql',           'MQL→SQL %',              b.get('mql_to_sql', 0.30) * 100, 1.0, '%.0f', True),
            ]
        elif li_fmt == 'Lead Gen Form':
            fields = [
                ('cpm',                  'CPM (€)',                 b['cpm'],                         0.1,  '%.2f', False),
                ('ctr',                  'CTR %',                   b['ctr'] * 100,                   0.01, '%.3f', True),
                ('form_completion_rate', 'Form Completion Rate %',  10.0,                             0.5,  '%.1f', True),
                ('lead_to_mql',          'Lead→MQL %',              60.0,                             1.0,  '%.0f', True),
                ('mql_to_sql',           'MQL→SQL %',               b.get('mql_to_sql', 0.30) * 100, 1.0,  '%.0f', True),
            ]
        else:  # Standard: Static / Video / Carousel
            fields = [
                ('cpm',       'CPM (€)',   b['cpm'],                              0.1,  '%.2f', False),
                ('ctr',       'CTR %',     b['ctr'] * 100,                        0.01, '%.3f', True),
                ('frequency', 'Frequency', b.get('frequency', 3.0),              0.5,  '%.1f', False),
            ]
            if goal in ('Traffic', 'Conversion'):
                fields.append(('click_to_session', 'Click→Session %',
                               b.get('click_to_session', 0.82) * 100, 1.0, '%.0f', True))
            if goal == 'Conversion':
                fields += [
                    ('conv_rate',   'Session→Lead %', b.get('conv_rate',   0.02) * 100, 0.1, '%.1f', True),
                    ('lead_to_mql', 'Lead→MQL %',     b.get('lead_to_mql', 0.20) * 100, 1.0, '%.0f', True),
                    ('mql_to_sql',  'MQL→SQL %',      b.get('mql_to_sql',  0.30) * 100, 1.0, '%.0f', True),
                ]

    else:  # Display
        fields = [
            ('cpm',       'CPM (€)',   b['cpm'],                              0.1,  '%.2f', False),
            ('ctr',       'CTR %',     b['ctr'] * 100,                        0.01, '%.3f', True),
            ('frequency', 'Frequency', b.get('frequency', 4.0),              0.5,  '%.1f', False),
        ]
        if goal in ('Traffic', 'Conversion'):
            fields.append(('click_to_session', 'Click→Session %',
                           b.get('click_to_session', 0.70) * 100, 1.0, '%.0f', True))
        if goal == 'Conversion':
            fields += [
                ('conv_rate',   'Session→Lead %', b.get('conv_rate',   0.02) * 100, 0.1, '%.1f', True),
                ('lead_to_mql', 'Lead→MQL %',     b.get('lead_to_mql', 0.20) * 100, 1.0, '%.0f', True),
                ('mql_to_sql',  'MQL→SQL %',      b.get('mql_to_sql',  0.30) * 100, 1.0, '%.0f', True),
            ]

    cols = st.columns(len(fields))
    raw = {}
    for i, (key, label, default, step, fmt, _) in enumerate(fields):
        raw[key] = cols[i].number_input(label, value=float(default), step=step, format=fmt,
                                        help=BENCH_HELP.get(key, ''),
                                        key=f'{key}_{mkt}_{ch}_{goal}_{sid}')

    bm = {}
    for key, _, _, _, _, is_pct in fields:
        bm[key] = raw[key] / 100.0 if is_pct else raw[key]

    # Store LinkedIn format so calc_row knows which funnel path to use
    if ch == 'LinkedIn':
        bm['_li_format'] = li_fmt

    # Carry over non-editable defaults
    if ch == 'YouTube' and goal != 'Awareness':
        bm['view_rate'] = b.get('view_rate', 0.31)
    if ch in ('YouTube', 'Display') and 'frequency' not in bm:
        bm['frequency'] = b.get('frequency', 4.0 if ch == 'Display' else 3.0)
    _li_no_freq = ('Sponsored Message / Conversational Ad', 'Conversation Ad',
                   'Lead Gen Form', 'Document Ad')
    if ch == 'LinkedIn' and li_fmt not in _li_no_freq:
        if 'frequency' not in bm:
            bm['frequency'] = b.get('frequency', 3.0)
    if 'click_to_session' not in bm:
        bm['click_to_session'] = b.get('click_to_session', 0.80)
    if 'conv_rate' not in bm:
        bm['conv_rate'] = b.get('conv_rate', 0.02)
    if 'lead_to_mql' not in bm:
        bm['lead_to_mql'] = b.get('lead_to_mql', 0.20)
    if 'mql_to_sql' not in bm:
        bm['mql_to_sql'] = b.get('mql_to_sql', 0.30)

    return bm


def _channel_budget_split(mkt, goal, goal_chs, mkt_budget, sid=0):
    """Render per-channel budget split for one market/goal pair. Returns {ch: budget}."""
    if len(goal_chs) == 1:
        return {goal_chs[0]: mkt_budget}

    st.markdown('**Channel budget split**')

    if len(goal_chs) == 2:
        ch_a, ch_b = goal_chs
        pct_a = st.slider(
            f'{ch_a}  ←——→  {ch_b}',
            min_value=0, max_value=100, value=50, step=5,
            key=f'split_{mkt}_{goal}_{sid}',
            format='%d%%',
        )
        pct_b = 100 - pct_a
        bud_a = mkt_budget * pct_a / 100
        bud_b = mkt_budget * pct_b / 100
        # Show the resulting amounts side by side
        ca, cb = st.columns(2)
        ca.metric(ch_a, f'€{bud_a:,.0f}', f'{pct_a}%')
        cb.metric(ch_b, f'€{bud_b:,.0f}', f'{pct_b}%')
        return {ch_a: bud_a, ch_b: bud_b}

    # 3 channels: individual % inputs that should sum to 100
    default_pct = round(100 / len(goal_chs), 1)
    pcts = {}
    cols = st.columns(len(goal_chs))
    for i, ch in enumerate(goal_chs):
        pcts[ch] = cols[i].number_input(
            f'{ch} %', min_value=0.0, max_value=100.0,
            value=default_pct, step=5.0, format='%.0f',
            key=f'split_{mkt}_{goal}_{ch}_{sid}'
        )
    pct_sum = sum(pcts.values())
    ch_budgets = {ch: mkt_budget * pcts[ch] / pct_sum for ch in goal_chs} if pct_sum > 0 else {ch: 0 for ch in goal_chs}
    if abs(pct_sum - 100) > 0.5:
        st.caption(f'Percentages sum to {pct_sum:.0f}% — normalising proportionally.')
    amt_cols = st.columns(len(goal_chs))
    for i, ch in enumerate(goal_chs):
        amt_cols[i].metric(ch, f'€{ch_budgets[ch]:,.0f}')
    return ch_budgets


def render_goal_section(mkt, goal, selected_channels, ch_budgets, periods, grand_totals, sid=0):
    """Render benchmarks + table + funnel for one market/goal combo."""
    bm_all = {}
    with st.expander('Edit benchmarks', expanded=True):
        for ch in selected_channels:
            st.markdown(f'**{ch}**')
            bm_all[ch] = benchmark_inputs(ch, mkt, goal, sid)

    for ch in selected_channels:
        bm = bm_all[ch]
        ch_bud = ch_budgets[ch]
        conv_rate = bm.get('conv_rate', 0.02)
        li_fmt = bm.get('_li_format') if ch == 'LinkedIn' else None

        df = build_table(periods, ch_bud, bm, goal, ch, conv_rate)

        col_t, col_f = st.columns([3, 2])
        with col_t:
            st.markdown(f'**{ch}** — €{ch_bud:,.0f}')
            st.dataframe(fmt_df(df, ch, goal, li_format=li_fmt), use_container_width=True, hide_index=True)
        with col_f:
            st.plotly_chart(
                make_funnel(df, goal, ch, f'{mkt} — {ch}', li_format=li_fmt),
                use_container_width=True,
                config={'displayModeBar': False},
                key=f'funnel_{mkt}_{ch}_{goal}_{sid}',
            )

        t_row = df[df['Period'] == 'TOTAL'].copy()
        t_row['Market'] = MARKET_LABELS[mkt]
        grand_totals[goal][ch].append(t_row)


# ── Save / Load helpers ───────────────────────────────────────────────────────

_SKIP_KEYS = {
    '_pending_load', 'FormSubmitter', '_uploader_v',
    # button / download widget keys — never safe to restore
    'dup_', 'remove_', 'tpl_apply_', 'grp_', 'eq_', 'cpm_eff_',
    'pin_', 'dl_gads_', 'dl_excel', 'btn_', 'preset_',
    'save_tpl_', 'del_tpl_',
    # derived / internal — recalculated on render, not worth saving
    'bud_mkt_', '_last_total_', 'li_fmt_prev_',
    # AI chat state — ephemeral, don't persist across plan loads
    'bench_chat', 'insights_chat', 'recs_chat',
    'insights_last', 'recs_last',
}

def _serialise_state():
    data = {'scenario_names': st.session_state.get('scenario_names', ['Scenario 1'])}
    # custom_templates is a dict of dicts — handle it explicitly
    if 'custom_templates' in st.session_state:
        data['custom_templates'] = st.session_state['custom_templates']
    for k, v in st.session_state.items():
        if any(k.startswith(s) for s in _SKIP_KEYS):
            continue
        if k == 'custom_templates':
            continue  # already handled above
        if isinstance(v, date):
            data[k] = {'__date__': v.isoformat()}
        elif isinstance(v, (str, int, float, bool, list)):
            data[k] = v
    return json.dumps(data, indent=2).encode('utf-8')


# ── Excel builder ─────────────────────────────────────────────────────────────

_NMQ_TEAL   = '2BB5A5'
_NMQ_BLACK  = '1A1A1A'
_NMQ_LIGHT  = 'F5F7F7'

def _xl_header(ws, row_num):
    fill = PatternFill('solid', fgColor=_NMQ_TEAL)
    font = Font(bold=True, color='FFFFFF', name='Calibri')
    for cell in ws[row_num]:
        if cell.value is not None:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal='center')


def _get_bm_ss(ch, mkt, goal, sid):
    """Read benchmark values from session state (mirrors benchmark_inputs without rendering)."""
    ss = st.session_state
    b  = BENCH[mkt][ch]

    def _pct(key, default):
        return ss.get(f'{key}_{mkt}_{ch}_{goal}_{sid}', default * 100) / 100

    bm = {}
    if ch == 'Search':
        bm['cpc']             = ss.get(f'cpc_{mkt}_{ch}_{goal}_{sid}', b['cpc'])
        bm['ctr']             = _pct('ctr',             b['ctr'])
        bm['click_to_session']= _pct('click_to_session',b.get('click_to_session', 0.85))
        bm['conv_rate']       = _pct('conv_rate',       b.get('conv_rate', 0.03))
    elif ch == 'YouTube':
        bm['cpm']             = ss.get(f'cpm_{mkt}_{ch}_{goal}_{sid}', b['cpm'])
        bm['ctr']             = _pct('ctr',             b['ctr'])
        bm['frequency']       = ss.get(f'frequency_{mkt}_{ch}_{goal}_{sid}', b.get('frequency', 3.0))
        bm['view_rate']       = (_pct('view_rate', b.get('view_rate', 0.31))
                                 if goal == 'Awareness' else b.get('view_rate', 0.31))
        bm['click_to_session']= _pct('click_to_session',b.get('click_to_session', 0.80))
        bm['conv_rate']       = _pct('conv_rate',       b.get('conv_rate', 0.02))
    elif ch == 'LinkedIn':
        li_fmt = ss.get(f'li_fmt_{mkt}_{goal}_{sid}', _LI_FORMATS[0])
        bm['_li_format'] = li_fmt
        if li_fmt == 'Sponsored Message / Conversational Ad':
            bm['cpm']       = ss.get(f'cpm_{mkt}_{ch}_{goal}_{sid}', b.get('cpm', 0.50))
            bm['open_rate'] = _pct('open_rate', 0.30)
            bm['ctr']       = _pct('ctr', b['ctr'])
            bm['conv_rate'] = _pct('conv_rate', 0.05)
        elif li_fmt == 'Conversation Ad':
            bm['cpm']       = ss.get(f'cpm_{mkt}_{ch}_{goal}_{sid}', b.get('cpm', 0.40))
            bm['open_rate'] = _pct('open_rate', 0.45)
            bm['ctr']       = _pct('ctr', b['ctr'])
            bm['conv_rate'] = _pct('conv_rate', 0.06)
        elif li_fmt == 'Document Ad':
            bm['cpm']                  = ss.get(f'cpm_{mkt}_{ch}_{goal}_{sid}', b['cpm'])
            bm['ctr']                  = _pct('ctr', b['ctr'])
            bm['form_completion_rate'] = _pct('form_completion_rate', 0.08)
        elif li_fmt == 'Lead Gen Form':
            bm['cpm']                  = ss.get(f'cpm_{mkt}_{ch}_{goal}_{sid}', b['cpm'])
            bm['ctr']                  = _pct('ctr', b['ctr'])
            bm['form_completion_rate'] = _pct('form_completion_rate', 0.10)
        else:  # Standard (Static / Video / Carousel)
            bm['cpm']             = ss.get(f'cpm_{mkt}_{ch}_{goal}_{sid}', b['cpm'])
            bm['ctr']             = _pct('ctr', b['ctr'])
            bm['frequency']       = ss.get(f'frequency_{mkt}_{ch}_{goal}_{sid}', b.get('frequency', 3.0))
            bm['click_to_session']= _pct('click_to_session', b.get('click_to_session', 0.82))
            bm['conv_rate']       = _pct('conv_rate', b.get('conv_rate', 0.02))
    else:  # Display
        bm['cpm']             = ss.get(f'cpm_{mkt}_{ch}_{goal}_{sid}', b['cpm'])
        bm['ctr']             = _pct('ctr',             b['ctr'])
        bm['frequency']       = ss.get(f'frequency_{mkt}_{ch}_{goal}_{sid}', b.get('frequency', 4.0))
        bm['click_to_session']= _pct('click_to_session', b.get('click_to_session', 0.70))
        bm['conv_rate']       = _pct('conv_rate',       b.get('conv_rate', 0.02))

    # MQL/SQL ratios — present for all channels when goal is Conversion
    bm['lead_to_mql'] = _pct('lead_to_mql', b.get('lead_to_mql', 0.20))
    bm['mql_to_sql']  = _pct('mql_to_sql',  b.get('mql_to_sql',  0.30))
    return bm


def _get_ch_budgets_ss(mkt, goal, goal_chs, mkt_budget, sid):
    """Read channel budget splits from session state without rendering UI."""
    ss = st.session_state
    if len(goal_chs) == 1:
        return {goal_chs[0]: mkt_budget}
    if len(goal_chs) == 2:
        ch_a, ch_b = goal_chs
        pct_a = ss.get(f'split_{mkt}_{goal}_{sid}', 50)
        return {ch_a: mkt_budget * pct_a / 100, ch_b: mkt_budget * (100 - pct_a) / 100}
    default_pct = round(100 / len(goal_chs), 1)
    pcts = {ch: ss.get(f'split_{mkt}_{goal}_{ch}_{sid}', default_pct) for ch in goal_chs}
    pct_sum = sum(pcts.values()) or 1
    return {ch: mkt_budget * pcts[ch] / pct_sum for ch in goal_chs}


def _build_excel_all(all_scenarios, scenario_ids, campaign_name, start_date, end_date, breakdown='Weekly'):
    """
    One workbook. One tab per scenario. All countries stacked vertically in the same tab
    (2-row gap between countries). A Scenario Totals section at the bottom sums across
    all countries via Excel SUM formulas referencing each country's TOTAL row.
    """
    from openpyxl.utils import get_column_letter

    # ── Colour palette ────────────────────────────────────────────────────────
    C_TITLE  = PatternFill('solid', fgColor='1F3864')   # dark navy — scenario title / totals
    C_MKT    = PatternFill('solid', fgColor='2BB5A5')   # NMQ teal  — country header
    C_GOAL   = PatternFill('solid', fgColor='1F3864')   # dark navy — goal section header
    C_CH     = PatternFill('solid', fgColor='2E75B6')   # medium blue — channel header
    C_ASSM_H = PatternFill('solid', fgColor='E2EFDA')   # light green — assumptions labels
    C_ASSM_V = PatternFill('solid', fgColor='FFFDE7')   # light yellow — editable values
    C_HDR    = PatternFill('solid', fgColor='4472C4')   # blue — column headers
    C_TOTAL  = PatternFill('solid', fgColor='CFE1F3')   # light blue — TOTAL rows
    C_STOT   = PatternFill('solid', fgColor='D6E4F0')   # slightly darker — scenario TOTAL
    C_ODD    = PatternFill('solid', fgColor='FFFFFF')
    C_EVEN   = PatternFill('solid', fgColor='EEF3FB')

    F_WHITE_B = Font(color='FFFFFF', bold=True)
    F_BOLD    = Font(bold=True)
    F_NORMAL  = Font()
    AC = Alignment(horizontal='center', vertical='center', wrap_text=True)
    AL = Alignment(horizontal='left',   vertical='center')

    _PCT_KEYS  = {'ctr', 'view_rate', 'click_to_session', 'conv_rate',
                  'lead_to_mql', 'mql_to_sql', 'cvr',
                  'open_rate', 'form_completion_rate'}
    _EUR_KEYS  = {'Budget', 'cpc', 'cpa', 'cpv', 'eff_cpm',
                  'cost_per_mql', 'cost_per_sql',
                  'cost_per_send', 'cost_per_open'}
    _ADDITIVE  = {'Budget', 'impressions', 'reach', 'views', 'clicks',
                  'sessions', 'conversions', 'mql', 'sql',
                  'sends', 'opens', 'cta_clicks', 'form_completions'}
    # Rate cols whose scenario-total value is derived from summed additive cols (standard path)
    _RATE_FROM = {
        'ctr':              ('clicks',      'impressions'),
        'click_to_session': ('sessions',    'clicks'),
        'conv_rate':        ('conversions', 'sessions'),
        'lead_to_mql':      ('mql',         'conversions'),
        'mql_to_sql':       ('sql',         'mql'),
    }
    # Format-specific rate derivations for scenario totals
    _RATE_FROM_SM = {
        'open_rate':  ('opens',       'sends'),
        'ctr':        ('cta_clicks',  'opens'),
        'conv_rate':  ('conversions', 'cta_clicks'),
        'lead_to_mql':('mql',         'conversions'),
        'mql_to_sql': ('sql',         'mql'),
    }
    _RATE_FROM_LGF = {
        'ctr':                  ('clicks',           'impressions'),
        'form_completion_rate': ('form_completions', 'clicks'),
        'lead_to_mql':          ('mql',              'form_completions'),
        'mql_to_sql':           ('sql',              'mql'),
    }
    _RATE_FROM_DA = {
        'ctr':                  ('clicks',           'impressions'),
        'form_completion_rate': ('form_completions', 'clicks'),
        'lead_to_mql':          ('mql',              'form_completions'),
        'mql_to_sql':           ('sql',              'mql'),
    }

    def _num_fmt(key):
        if key in _PCT_KEYS: return '0.00%'
        if key in _EUR_KEYS: return '#,##0.00'
        return '#,##0'

    _BM_PARAMS = {
        ('YouTube',       'Awareness'):   ['cpm', 'view_rate', 'ctr', 'frequency'],
        ('YouTube',       'Traffic'):     ['cpm', 'ctr', 'click_to_session'],
        ('YouTube',       'Conversion'):  ['cpm', 'ctr', 'click_to_session', 'conv_rate', 'lead_to_mql', 'mql_to_sql'],
        ('LinkedIn',      'Awareness'):   ['cpm', 'ctr', 'frequency'],
        ('LinkedIn',      'Traffic'):     ['cpm', 'ctr', 'click_to_session'],
        ('LinkedIn',      'Conversion'):  ['cpm', 'ctr', 'click_to_session', 'conv_rate', 'lead_to_mql', 'mql_to_sql'],
        ('LinkedIn_SM',   'Awareness'):   ['cpm', 'open_rate', 'ctr'],
        ('LinkedIn_SM',   'Traffic'):     ['cpm', 'open_rate', 'ctr'],
        ('LinkedIn_SM',   'Conversion'):  ['cpm', 'open_rate', 'ctr', 'conv_rate', 'lead_to_mql', 'mql_to_sql'],
        ('LinkedIn_CA',   'Awareness'):   ['cpm', 'open_rate', 'ctr'],
        ('LinkedIn_CA',   'Traffic'):     ['cpm', 'open_rate', 'ctr'],
        ('LinkedIn_CA',   'Conversion'):  ['cpm', 'open_rate', 'ctr', 'conv_rate', 'lead_to_mql', 'mql_to_sql'],
        ('LinkedIn_DA',   'Awareness'):   ['cpm', 'ctr', 'form_completion_rate'],
        ('LinkedIn_DA',   'Traffic'):     ['cpm', 'ctr', 'form_completion_rate'],
        ('LinkedIn_DA',   'Conversion'):  ['cpm', 'ctr', 'form_completion_rate', 'lead_to_mql', 'mql_to_sql'],
        ('LinkedIn_LGF',  'Awareness'):   ['cpm', 'ctr', 'form_completion_rate'],
        ('LinkedIn_LGF',  'Traffic'):     ['cpm', 'ctr', 'form_completion_rate'],
        ('LinkedIn_LGF',  'Conversion'):  ['cpm', 'ctr', 'form_completion_rate', 'lead_to_mql', 'mql_to_sql'],
        ('Search',        'Awareness'):   ['cpc', 'ctr'],
        ('Search',        'Traffic'):     ['cpc', 'ctr', 'click_to_session'],
        ('Search',        'Conversion'):  ['cpc', 'ctr', 'click_to_session', 'conv_rate', 'lead_to_mql', 'mql_to_sql'],
        ('Display',       'Awareness'):   ['cpm', 'ctr', 'frequency'],
        ('Display',       'Traffic'):     ['cpm', 'ctr', 'click_to_session'],
        ('Display',       'Conversion'):  ['cpm', 'ctr', 'click_to_session', 'conv_rate', 'lead_to_mql', 'mql_to_sql'],
    }
    _BM_LABELS = {
        'cpm': 'CPM / Cost per Send (€)', 'cpc': 'CPC (€)', 'view_rate': 'View Rate',
        'ctr': 'CTR / CTA Click Rate', 'frequency': 'Frequency',
        'click_to_session': 'Click→Session', 'open_rate': 'Open Rate',
        'form_completion_rate': 'Form Completion Rate',
        'conv_rate': 'Session→Lead % / CTA→Lead %',
        'lead_to_mql': 'Lead→MQL %', 'mql_to_sql': 'MQL→SQL %',
    }
    _BM_NUM_FMT = {
        'cpm': '#,##0.0000', 'cpc': '#,##0.00', 'view_rate': '0.00%',
        'ctr': '0.00%', 'frequency': '0.0', 'click_to_session': '0%',
        'open_rate': '0.0%', 'form_completion_rate': '0.0%',
        'conv_rate': '0.00%', 'lead_to_mql': '0%', 'mql_to_sql': '0%',
    }

    def _formula(key, r, col_map, bm_map, ch, li_fmt=None):
        L = col_map
        def ref(k):   return f"{L.get(k,'A')}{r}"
        def bm(k):    return bm_map.get(k, '0')
        def ie(f):    return f'=IFERROR({f},"")'

        if key == 'Budget':       return None

        # ── LinkedIn Sponsored Message / Conversation Ad ──────────────────────
        if ch == 'LinkedIn' and li_fmt in (
                'Sponsored Message / Conversational Ad', 'Conversation Ad'):
            if key == 'sends':            return ie(f"{ref('Budget')}/{bm('cpm')}")
            if key == 'cost_per_send':    return ie(f"{ref('Budget')}/{ref('sends')}")
            if key == 'opens':            return ie(f"{ref('sends')}*{bm('open_rate')}")
            if key == 'cost_per_open':    return ie(f"{ref('Budget')}/{ref('opens')}")
            if key == 'open_rate':        return f"={bm('open_rate')}"
            if key == 'cta_clicks':       return ie(f"{ref('opens')}*{bm('ctr')}")
            if key == 'ctr':              return ie(f"{ref('cta_clicks')}/{ref('opens')}")
            if key == 'conv_rate':        return f"={bm('conv_rate')}"
            if key == 'conversions':      return ie(f"{ref('cta_clicks')}*{bm('conv_rate')}")
            if key == 'cpa':              return ie(f"{ref('Budget')}/{ref('conversions')}")
            if key == 'lead_to_mql':      return f"={bm('lead_to_mql')}"
            if key == 'mql':              return ie(f"{ref('conversions')}*{bm('lead_to_mql')}")
            if key == 'cost_per_mql':     return ie(f"{ref('Budget')}/{ref('mql')}")
            if key == 'mql_to_sql':       return f"={bm('mql_to_sql')}"
            if key == 'sql':              return ie(f"{ref('mql')}*{bm('mql_to_sql')}")
            if key == 'cost_per_sql':     return ie(f"{ref('Budget')}/{ref('sql')}")
            return None

        # ── LinkedIn Lead Gen Form + Document Ad (same calc path) ────────────
        if ch == 'LinkedIn' and li_fmt in ('Lead Gen Form', 'Document Ad'):
            if key == 'impressions':          return ie(f"{ref('Budget')}/{bm('cpm')}*1000")
            if key == 'eff_cpm':              return ie(f"{ref('Budget')}/{ref('impressions')}*1000")
            if key == 'clicks':               return ie(f"{ref('impressions')}*{bm('ctr')}")
            if key == 'ctr':                  return ie(f"{ref('clicks')}/{ref('impressions')}")
            if key == 'form_completions':     return ie(f"{ref('clicks')}*{bm('form_completion_rate')}")
            if key == 'form_completion_rate': return f"={bm('form_completion_rate')}"
            if key == 'lead_to_mql':          return f"={bm('lead_to_mql')}"
            if key == 'mql':                  return ie(f"{ref('form_completions')}*{bm('lead_to_mql')}")
            if key == 'cost_per_mql':         return ie(f"{ref('Budget')}/{ref('mql')}")
            if key == 'mql_to_sql':           return f"={bm('mql_to_sql')}"
            if key == 'sql':                  return ie(f"{ref('mql')}*{bm('mql_to_sql')}")
            if key == 'cost_per_sql':         return ie(f"{ref('Budget')}/{ref('sql')}")
            return None

        # ── Standard path (all other channels + LinkedIn Standard) ────────────
        if key == 'impressions':
            return ie(f"{ref('clicks')}/{bm('ctr')}") if ch == 'Search' \
                   else ie(f"{ref('Budget')}/{bm('cpm')}*1000")
        if key == 'reach':        return ie(f"{ref('impressions')}/{bm('frequency')}")
        if key == 'views':        return ie(f"{ref('impressions')}*{bm('view_rate')}")
        if key == 'clicks':
            return ie(f"{ref('Budget')}/{bm('cpc')}") if ch == 'Search' \
                   else ie(f"{ref('impressions')}*{bm('ctr')}")
        if key == 'ctr':          return ie(f"{ref('clicks')}/{ref('impressions')}")
        if key == 'cpc':          return ie(f"{ref('Budget')}/{ref('clicks')}")
        if key == 'cpv':          return ie(f"{ref('Budget')}/{ref('views')}")
        if key == 'eff_cpm':      return ie(f"{ref('Budget')}/{ref('impressions')}*1000")
        if key == 'click_to_session': return f"={bm('click_to_session')}"
        if key == 'sessions':     return ie(f"{ref('clicks')}*{bm('click_to_session')}")
        if key == 'conv_rate':    return f"={bm('conv_rate')}"
        if key == 'conversions':  return ie(f"{ref('sessions')}*{bm('conv_rate')}")
        if key == 'cpa':          return ie(f"{ref('Budget')}/{ref('conversions')}")
        if key == 'lead_to_mql':  return f"={bm('lead_to_mql')}"
        if key == 'mql':          return ie(f"{ref('conversions')}*{bm('lead_to_mql')}")
        if key == 'cost_per_mql': return ie(f"{ref('Budget')}/{ref('mql')}")
        if key == 'mql_to_sql':   return f"={bm('mql_to_sql')}"
        if key == 'sql':          return ie(f"{ref('mql')}*{bm('mql_to_sql')}")
        if key == 'cost_per_sql': return ie(f"{ref('Budget')}/{ref('sql')}")
        return None

    def _total_formula(key, r, col_map, bm_map, first_r, last_r, ch, li_fmt=None):
        L = col_map
        if key in _ADDITIVE:
            return f"=SUM({L.get(key,'A')}{first_r}:{L.get(key,'A')}{last_r})"
        # Rate fields that should reference the assumption cell, not be re-derived
        _rate_keys_sm  = ('open_rate', 'conv_rate', 'lead_to_mql', 'mql_to_sql')
        _rate_keys_lgf = ('form_completion_rate', 'lead_to_mql', 'mql_to_sql')
        _rate_keys_std = ('click_to_session', 'conv_rate', 'lead_to_mql', 'mql_to_sql')
        if ch == 'LinkedIn' and li_fmt in (
                'Sponsored Message / Conversational Ad', 'Conversation Ad'):
            if key in _rate_keys_sm:
                return f"={bm_map.get(key, '0')}"
        elif ch == 'LinkedIn' and li_fmt in ('Document Ad', 'Lead Gen Form'):
            if key in _rate_keys_lgf:
                return f"={bm_map.get(key, '0')}"
        else:
            if key in _rate_keys_std:
                return f"={bm_map.get(key, '0')}"
        return _formula(key, r, col_map, bm_map, ch, li_fmt=li_fmt)

    # ── Setup ─────────────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    periods    = generate_periods(start_date, end_date, breakdown)
    total_days = sum(p['days'] for p in periods) or 1
    flight     = f"{start_date.strftime('%b %d, %Y')} – {end_date.strftime('%b %d, %Y')}"

    # ── One tab per scenario ──────────────────────────────────────────────────
    for s, sid in zip(all_scenarios, scenario_ids):
        tab_name = s['name'][:31]
        ws = wb.create_sheet(title=tab_name)

        max_data_cols = max(
            (len(PHASE_COLS.get((ch, goal), [])) for goal, chs in s['goal_channels'].items()
             for ch in chs), default=8
        )
        n_cols_total = 1 + max_data_cols

        def _merge(r, fill, font, txt, height=18, align=AL):
            c = ws.cell(row=r, column=1, value=txt)
            c.fill = fill; c.font = font; c.alignment = align
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n_cols_total)
            ws.row_dimensions[r].height = height

        # ── Scenario title rows ───────────────────────────────────────────────
        row = 1
        _merge(row, C_TITLE, Font(color='FFFFFF', bold=True, size=13),
               f"{s['name']}  —  {campaign_name}", height=22)
        row += 1
        _merge(row, C_TITLE, Font(color='FFFFFF', size=10),
               f"Total Budget: €{s['market_budgets'].get(s['selected_markets'][0] if s['selected_markets'] else '', 0):,.0f}"
               f"  (see per-country below)  |  {flight}  |  {breakdown}", height=15)
        row += 2  # title rows + blank

        # ── Track each country's TOTAL row position for scenario-totals SUM ──
        # total_tracker[(goal, ch)] = {'rows': [row_nums], 'col_keys': [...], 'col_map': {...}}
        total_tracker = {}

        # ── Countries — stacked vertically ───────────────────────────────────
        for mkt_idx, mkt in enumerate(s['selected_markets']):
            mkt_label  = MARKET_LABELS.get(mkt, mkt)
            mkt_budget = s['market_budgets'][mkt]
            total_budget_all = sum(s['market_budgets'].values()) or 1
            mkt_pct = mkt_budget / total_budget_all * 100

            # Country header (NMQ teal)
            _merge(row, C_MKT, F_WHITE_B,
                   f"{mkt_label}  —  €{mkt_budget:,.0f}  ({mkt_pct:.0f}%)", height=20,
                   align=Alignment(horizontal='center', vertical='center'))
            row += 1

            for goal in ['Awareness', 'Traffic', 'Conversion']:
                if goal not in s['goal_channels']:
                    continue
                goal_chs   = s['goal_channels'][goal]
                ch_budgets = _get_ch_budgets_ss(mkt, goal, goal_chs, mkt_budget, sid)

                # Goal header
                _merge(row, C_GOAL, F_WHITE_B, goal.upper(), height=18,
                       align=Alignment(horizontal='center', vertical='center'))
                row += 1

                for ch in goal_chs:
                    ch_bud    = ch_budgets.get(ch, 0)
                    bm_data   = _get_bm_ss(ch, mkt, goal, sid)
                    li_fmt    = bm_data.get('_li_format') if ch == 'LinkedIn' else None

                    # Resolve format-specific PHASE_COLS key for LinkedIn
                    if ch == 'LinkedIn' and li_fmt == 'Sponsored Message / Conversational Ad':
                        phase_key = ('LinkedIn_SM', goal)
                        bm_params_key = ('LinkedIn_SM', goal)
                    elif ch == 'LinkedIn' and li_fmt == 'Conversation Ad':
                        phase_key = ('LinkedIn_CA', goal)
                        bm_params_key = ('LinkedIn_CA', goal)
                    elif ch == 'LinkedIn' and li_fmt == 'Document Ad':
                        phase_key = ('LinkedIn_DA', goal)
                        bm_params_key = ('LinkedIn_DA', goal)
                    elif ch == 'LinkedIn' and li_fmt == 'Lead Gen Form':
                        phase_key = ('LinkedIn_LGF', goal)
                        bm_params_key = ('LinkedIn_LGF', goal)
                    else:
                        phase_key = (ch, goal)
                        bm_params_key = (ch, goal)

                    col_keys  = PHASE_COLS.get(phase_key, list(_TRAFFIC_COLS))
                    n_ch_cols = len(col_keys) + 1
                    col_map   = {key: get_column_letter(2 + i) for i, key in enumerate(col_keys)}
                    daily_bud = ch_bud / total_days if total_days else 0

                    # Channel sub-header
                    ch_label = f'{ch} ({li_fmt})' if li_fmt else ch
                    c = ws.cell(row=row, column=1,
                                value=f"{ch_label}     Daily Budget: €{daily_bud:,.2f} / day")
                    c.fill = C_CH; c.font = F_WHITE_B
                    c.alignment = Alignment(horizontal='center', vertical='center')
                    ws.merge_cells(start_row=row, start_column=1,
                                   end_row=row, end_column=n_ch_cols)
                    ws.row_dimensions[row].height = 18
                    row += 1

                    # Assumptions block
                    bm_params = _BM_PARAMS.get(bm_params_key, [])
                    c = ws.cell(row=row, column=1,
                                value='ASSUMPTIONS  —  edit the yellow cells to recalculate the whole table')
                    c.fill = C_ASSM_H; c.font = Font(italic=True, size=9, color='375623')
                    c.alignment = AL
                    ws.merge_cells(start_row=row, start_column=1,
                                   end_row=row, end_column=n_ch_cols)
                    ws.row_dimensions[row].height = 14
                    row += 1

                    bm_col_map = {p: get_column_letter(2 + i) for i, p in enumerate(bm_params)}
                    for i, param in enumerate(bm_params):
                        c = ws.cell(row=row, column=2 + i, value=_BM_LABELS.get(param, param))
                        c.fill = C_ASSM_H; c.font = Font(bold=True, size=9, color='375623')
                        c.alignment = AC
                    ws.row_dimensions[row].height = 14
                    row += 1

                    bm_val_row = row
                    bm_map     = {p: f'${bm_col_map[p]}${bm_val_row}' for p in bm_params}
                    for i, param in enumerate(bm_params):
                        c = ws.cell(row=row, column=2 + i, value=bm_data.get(param, 0))
                        c.fill = C_ASSM_V; c.font = Font(bold=True, size=10); c.alignment = AC
                        c.number_format = _BM_NUM_FMT.get(param, 'General')
                    ws.row_dimensions[row].height = 18
                    row += 2  # values + blank

                    # Column headers
                    c = ws.cell(row=row, column=1, value='PERIOD')
                    c.fill = C_HDR; c.font = F_WHITE_B; c.alignment = AC
                    for ci, key in enumerate(col_keys, 2):
                        c = ws.cell(row=row, column=ci,
                                    value=COL_FMT[key][0] if key in COL_FMT else key)
                        c.fill = C_HDR; c.font = F_WHITE_B; c.alignment = AC
                    ws.row_dimensions[row].height = 28
                    row += 1

                    # Period data rows
                    first_data_row = row
                    for r_idx, p in enumerate(periods):
                        bud  = ch_bud * p['days'] / total_days if total_days else 0
                        fill = C_ODD if r_idx % 2 == 0 else C_EVEN
                        c = ws.cell(row=row, column=1, value=str(p['label']))
                        c.fill = fill; c.font = F_BOLD; c.alignment = AC
                        for ci, key in enumerate(col_keys, 2):
                            if key == 'Budget':
                                c = ws.cell(row=row, column=ci, value=bud)
                            else:
                                f = _formula(key, row, col_map, bm_map, ch, li_fmt=li_fmt)
                                c = ws.cell(row=row, column=ci, value=f if f is not None else 0)
                            c.fill = fill; c.font = F_NORMAL; c.alignment = AC
                            c.number_format = _num_fmt(key)
                        row += 1
                    last_data_row = row - 1

                    # TOTAL row — track its row number for scenario-totals formulas
                    total_row_num = row
                    c = ws.cell(row=row, column=1, value='TOTAL')
                    c.fill = C_TOTAL; c.font = F_BOLD; c.alignment = AC
                    for ci, key in enumerate(col_keys, 2):
                        f = _total_formula(key, row, col_map, bm_map,
                                           first_data_row, last_data_row, ch, li_fmt=li_fmt)
                        c = ws.cell(row=row, column=ci, value=f if f is not None else 0)
                        c.fill = C_TOTAL; c.font = F_BOLD; c.alignment = AC
                        c.number_format = _num_fmt(key)
                    ws.row_dimensions[row].height = 15
                    row += 2

                    k = (goal, ch)
                    if k not in total_tracker:
                        total_tracker[k] = {'rows': [], 'col_keys': col_keys, 'col_map': col_map,
                                            'li_fmt': li_fmt}
                    total_tracker[k]['rows'].append(total_row_num)

                row += 1  # blank between goals

            # 2-row gap between countries (as requested)
            if mkt_idx < len(s['selected_markets']) - 1:
                row += 2

        # ── Scenario Totals ───────────────────────────────────────────────────
        row += 2
        _merge(row, C_TITLE, Font(color='FFFFFF', bold=True, size=13),
               f"SCENARIO TOTALS  —  {s['name']}  —  {len(s['selected_markets'])} countr{'y' if len(s['selected_markets'])==1 else 'ies'}",
               height=22, align=Alignment(horizontal='center', vertical='center'))
        row += 2

        for goal in ['Awareness', 'Traffic', 'Conversion']:
            if goal not in s['goal_channels']:
                continue
            goal_chs = s['goal_channels'][goal]

            _merge(row, C_GOAL, F_WHITE_B, goal.upper(), height=18,
                   align=Alignment(horizontal='center', vertical='center'))
            row += 1

            for ch in goal_chs:
                k = (goal, ch)
                if k not in total_tracker:
                    continue
                tr      = total_tracker[k]
                ckeys   = tr['col_keys']
                cmap    = tr['col_map']
                trows   = tr['rows']
                li_fmt  = tr.get('li_fmt')
                n_ch_cols = len(ckeys) + 1

                # Choose the right rate derivation map for this format
                if ch == 'LinkedIn' and li_fmt in (
                        'Sponsored Message / Conversational Ad', 'Conversation Ad'):
                    rate_from = _RATE_FROM_SM
                elif ch == 'LinkedIn' and li_fmt == 'Document Ad':
                    rate_from = _RATE_FROM_DA
                elif ch == 'LinkedIn' and li_fmt == 'Lead Gen Form':
                    rate_from = _RATE_FROM_LGF
                else:
                    rate_from = _RATE_FROM

                # Channel header (no daily budget — this is the aggregate)
                ch_label = f'{ch} ({li_fmt})' if li_fmt else ch
                c = ws.cell(row=row, column=1, value=f"{ch_label}  —  All Countries Combined")
                c.fill = C_CH; c.font = F_WHITE_B
                c.alignment = Alignment(horizontal='center', vertical='center')
                ws.merge_cells(start_row=row, start_column=1,
                               end_row=row, end_column=n_ch_cols)
                ws.row_dimensions[row].height = 18
                row += 1

                # Column headers
                c = ws.cell(row=row, column=1, value='')
                c.fill = C_HDR; c.alignment = AC
                for ci, key in enumerate(ckeys, 2):
                    c = ws.cell(row=row, column=ci,
                                value=COL_FMT[key][0] if key in COL_FMT else key)
                    c.fill = C_HDR; c.font = F_WHITE_B; c.alignment = AC
                ws.row_dimensions[row].height = 28
                row += 1

                # Scenario-total row
                c = ws.cell(row=row, column=1, value='ALL COUNTRIES')
                c.fill = C_STOT; c.font = F_BOLD; c.alignment = AC
                for ci, key in enumerate(ckeys, 2):
                    col_letter = cmap.get(key, 'B')
                    if key in _ADDITIVE:
                        refs = ', '.join(f'{col_letter}{r}' for r in trows)
                        f = f'=IFERROR(SUM({refs}),"")'
                    elif key in rate_from:
                        nk, dk = rate_from[key]
                        if nk in cmap and dk in cmap:
                            f = f'=IFERROR({cmap[nk]}{row}/{cmap[dk]}{row},"")'
                        else:
                            f = ''
                    else:
                        # derived cost fields — reference same row's additive cells
                        f = _formula(key, row, cmap, {}, ch, li_fmt=li_fmt)
                    c = ws.cell(row=row, column=ci, value=f if f else '')
                    c.fill = C_STOT; c.font = F_BOLD; c.alignment = AC
                    c.number_format = _num_fmt(key)
                ws.row_dimensions[row].height = 18
                row += 2

            row += 1  # blank between goals

        # ── Column widths ─────────────────────────────────────────────────────
        ws.column_dimensions['A'].width = 18
        for i in range(2, n_cols_total + 1):
            ws.column_dimensions[get_column_letter(i)].width = 12

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ── Google Ads CSV builder ────────────────────────────────────────────────────

_MARKET_LANGUAGE = {
    'AT': 'German', 'BE': 'French', 'BG': 'Bulgarian', 'HR': 'Croatian',
    'CY': 'Greek', 'CZ': 'Czech', 'DK': 'Danish', 'EE': 'Estonian',
    'FI': 'Finnish', 'FR': 'French', 'DE': 'German', 'GR': 'Greek',
    'HU': 'Hungarian', 'IE': 'English', 'IT': 'Italian', 'LV': 'Latvian',
    'LT': 'Lithuanian', 'LU': 'French', 'MT': 'English', 'NL': 'Dutch',
    'NO': 'Norwegian', 'PL': 'Polish', 'PT': 'Portuguese', 'RO': 'Romanian',
    'SK': 'Slovak', 'SI': 'Slovenian', 'ES': 'Spanish', 'SE': 'Swedish',
    'CH': 'German', 'UK': 'English',
}
_GADS_TYPE = {'YouTube': 'Video', 'Search': 'Search', 'Display': 'Display'}
_GADS_BID = {
    ('YouTube',  'Awareness'):  'Target CPM',
    ('YouTube',  'Traffic'):    'Maximize conversions',
    ('YouTube',  'Conversion'): 'Target CPA',
    ('Search',   'Awareness'):  'Manual CPC',
    ('Search',   'Traffic'):    'Maximize clicks',
    ('Search',   'Conversion'): 'Target CPA',
    ('Display',  'Awareness'):  'Target CPM',
    ('Display',  'Traffic'):    'Maximize clicks',
    ('Display',  'Conversion'): 'Target CPA',
}
_GADS_COLS = [
    'Campaign', 'Campaign Status', 'Campaign Type',
    'Budget', 'Budget Type',
    'Bid Strategy Type', 'Target CPA', 'Target CPM',
    'Start Date', 'End Date',
    'Location', 'Location Type',
    'Language', 'Language Type',
    'Ad Group', 'Ad Group Status', 'Default Max. CPC',
]


def _build_gads_csv_scenario(s_data, sid, campaign_name, start_date, end_date):
    """Build a Google Ads Editor CSV directly from a rendered scenario's data."""
    ss   = st.session_state
    days = max((end_date - start_date).days + 1, 1)
    s_dt = start_date.strftime('%m/%d/%Y')
    e_dt = end_date.strftime('%m/%d/%Y')
    rows = []

    def _empty():
        return {c: '' for c in _GADS_COLS}

    for mkt in s_data['selected_markets']:
        mkt_budget = s_data['market_budgets'][mkt]
        for goal, channels in s_data['goal_channels'].items():
            gads_chs = [ch for ch in channels if ch in _GADS_TYPE]
            if not gads_chs:
                continue

            if len(gads_chs) == 1:
                ch_buds = {gads_chs[0]: mkt_budget}
            elif len(gads_chs) == 2:
                pct_a = ss.get(f'split_{mkt}_{goal}_{sid}', 50)
                ch_buds = {
                    gads_chs[0]: mkt_budget * pct_a / 100,
                    gads_chs[1]: mkt_budget * (100 - pct_a) / 100,
                }
            else:
                pcts = {ch: ss.get(f'split_{mkt}_{goal}_{ch}_{sid}', 100 / len(gads_chs)) for ch in gads_chs}
                total = sum(pcts.values()) or 1
                ch_buds = {ch: mkt_budget * pcts[ch] / total for ch in gads_chs}

            for ch in gads_chs:
                daily      = round(ch_buds[ch] / days, 2)
                name       = f'{campaign_name}_{mkt}_{goal}_{ch}'
                bid        = _GADS_BID.get((ch, goal), 'Manual CPC')
                target_cpm = ss.get(f'cpm_{mkt}_{ch}_{goal}_{sid}', '') if ch in ('YouTube', 'Display') else ''
                default_cpc = ss.get(f'cpc_{mkt}_{ch}_{goal}_{sid}', '') if ch == 'Search' else ''

                r = _empty()
                r.update({'Campaign': name, 'Campaign Status': 'Enabled',
                          'Campaign Type': _GADS_TYPE[ch], 'Budget': daily,
                          'Budget Type': 'Daily', 'Bid Strategy Type': bid,
                          'Target CPM': target_cpm, 'Start Date': s_dt, 'End Date': e_dt})
                rows.append(r)

                r = _empty()
                r.update({'Campaign': name, 'Location': MARKET_LABELS[mkt], 'Location Type': 'Location'})
                rows.append(r)

                r = _empty()
                r.update({'Campaign': name, 'Language': _MARKET_LANGUAGE.get(mkt, 'English'), 'Language Type': 'Language'})
                rows.append(r)

                r = _empty()
                r.update({'Campaign': name, 'Ad Group': f'{name}_AdGroup_01',
                          'Ad Group Status': 'Enabled', 'Default Max. CPC': default_cpc})
                rows.append(r)

    buf = io.StringIO()
    pd.DataFrame(rows, columns=_GADS_COLS).to_csv(buf, index=False)
    return buf.getvalue().encode('utf-8')


# ── Step label helper ─────────────────────────────────────────────────────────
def _step(n, label):
    return (
        f'<p style="margin:6px 0 2px 0;font-weight:600;font-size:0.88rem;color:#1A1A1A">'
        f'<span style="background:#2BB5A5;color:white;border-radius:50%;'
        f'width:18px;height:18px;display:inline-flex;align-items:center;justify-content:center;'
        f'font-size:0.68rem;font-weight:700;margin-right:6px;flex-shrink:0">{n}</span>'
        f'{label}</p>'
    )


# Apply any pending plan load before widgets render (backup for app.py handler)
if '_pending_load' in st.session_state:
    _load_data = st.session_state.pop('_pending_load')
    for _k, _v in _load_data.items():
        if any(_k.startswith(s) for s in _SKIP_KEYS):
            continue
        if isinstance(_v, dict) and '__date__' in _v:
            st.session_state[_k] = date.fromisoformat(_v['__date__'])
        else:
            st.session_state[_k] = _v


# ── Sidebar ───────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown(_step(1, 'Campaign Name'), unsafe_allow_html=True)
    campaign_name = st.text_input('Campaign Name', value='Campaign — 2026',
                                  key='campaign_name', label_visibility='collapsed')

    st.markdown(_step(2, 'Audience Type'), unsafe_allow_html=True)
    audience_type = st.radio('Audience', ['B2B', 'B2C'], horizontal=True,
                             label_visibility='collapsed', key='audience_type')

    st.markdown(_step(3, 'Industry'), unsafe_allow_html=True)
    INDUSTRIES_SB = [
        'Logistics, Supply Chain & Transportation',
        'Industrial, Manufacturing & Materials',
        'Enterprise SaaS, Technology & Platforms',
        'Financial Services, Fintech & Insurance',
        'Healthcare, Pharma & Life Sciences',
        'Consumer Brands & Retail',
        'Mobility, Travel & Automotive',
        'Information Services & Professional Platforms',
    ]
    industry = st.selectbox('Industry', INDUSTRIES_SB, label_visibility='collapsed', key='industry')

    st.markdown('---')
    st.markdown(_step(4, 'Flight Dates'), unsafe_allow_html=True)
    col_s, col_e = st.columns(2)
    start_date = col_s.date_input('Start', value=date(2026, 5, 12), key='start_date')
    end_date   = col_e.date_input('End',   value=date(2026, 6, 29), key='end_date')

    if end_date <= start_date:
        st.error('End must be after Start.')
        st.stop()

    st.caption(f'{(end_date - start_date).days + 1} days total')

    st.markdown(_step(5, 'Period Breakdown'), unsafe_allow_html=True)
    breakdown = st.selectbox('Breakdown', ['Daily', 'Weekly', 'Bi-Weekly', 'Monthly'],
                             index=1, label_visibility='collapsed', key='breakdown')

    # ── Save / Load ───────────────────────────────────────────────────────────
    st.markdown('---')
    st.markdown('**💾 Save / Load Plan**')
    st.download_button(
        label='Download plan (.json)',
        data=_serialise_state(),
        file_name=f'{campaign_name.replace(" ", "_")}_plan.json',
        mime='application/json',
        use_container_width=True,
    )
    if '_uploader_v' not in st.session_state:
        st.session_state['_uploader_v'] = 0
    uploaded = st.file_uploader('Load plan (.json)', type='json',
                                label_visibility='collapsed',
                                key=f'_plan_uploader_{st.session_state["_uploader_v"]}')
    if uploaded is not None:
        try:
            payload = json.loads(uploaded.read().decode('utf-8'))
            st.session_state['_pending_load'] = payload
            st.session_state['_uploader_v'] += 1  # new key on next render = clears the uploader
            st.rerun()
        except Exception as e:
            st.error(f'Could not load plan: {e}')


# ── Auto-save warning ─────────────────────────────────────────────────────────
_components.html("""
<script>
window.addEventListener('beforeunload', function(e) {
    e.preventDefault();
    e.returnValue = 'Download your plan JSON from the sidebar before leaving — unsaved changes will be lost.';
});
</script>
""", height=0)

# ── Main ──────────────────────────────────────────────────────────────────────
st.markdown(
    f'<h2 style="margin:0 0 2px 0;font-family:Inter,sans-serif;font-size:1.5rem;color:#1A1A1A">'
    f'{campaign_name}</h2>',
    unsafe_allow_html=True,
)
st.caption(f'{start_date.strftime("%b %d, %Y")} – {end_date.strftime("%b %d, %Y")}  ·  {breakdown}  ·  {audience_type}  ·  {industry}')
st.divider()

# Scenario management via session state
if 'scenario_names' not in st.session_state:
    st.session_state.scenario_names = ['Scenario 1']
if 'custom_templates' not in st.session_state:
    st.session_state['custom_templates'] = {}

add_col, _ = st.columns([1, 6])
if add_col.button('＋ Add Scenario'):
    n = len(st.session_state.scenario_names) + 1
    st.session_state.scenario_names.append(f'Scenario {n}')

periods = generate_periods(start_date, end_date, breakdown)

_has_compare = len(st.session_state.scenario_names) >= 2
_tab_labels = (['⚖ Compare'] if _has_compare else []) + st.session_state.scenario_names
_all_tabs = st.tabs(_tab_labels)
compare_tab = _all_tabs[0] if _has_compare else None
scenario_tabs = _all_tabs[1:] if _has_compare else _all_tabs
all_scenarios_data = []  # collected for AI section and Compare tab


def _render_scenario(sid):
    """Render per-scenario config (goals, channels, markets, budget, split) then tables and funnels."""

    # ── Header: rename + duplicate + remove ───────────────────────────────────
    n_col, dup_col, rem_col = st.columns([6, 1, 1])
    new_name = n_col.text_input('Scenario name', value=st.session_state.scenario_names[sid],
                                key=f'rename_{sid}', label_visibility='collapsed',
                                placeholder='Scenario name…')
    if new_name and new_name != st.session_state.scenario_names[sid]:
        st.session_state.scenario_names[sid] = new_name
        st.session_state['_pending_dup_rerun'] = True
    if dup_col.button('⧉ Dup', key=f'dup_{sid}', use_container_width=True, help='Duplicate this scenario'):
        _duplicate_scenario(sid)
    if rem_col.button('✕ Remove', key=f'remove_{sid}', use_container_width=True,
                      disabled=len(st.session_state.scenario_names) <= 1,
                      help='Remove this scenario'):
        st.session_state.scenario_names.pop(sid)
        st.session_state['_pending_dup_rerun'] = True

    # ── Completion status ─────────────────────────────────────────────────────
    _scenario_status(sid)

    # ── Templates ─────────────────────────────────────────────────────────────
    with st.expander('📁 Templates', expanded=False):
        custom_tpls = st.session_state.get('custom_templates', {})

        # ── Combined dropdown (built-in + saved) ──────────────────────────────
        bi_options   = list(PLAN_TEMPLATES.keys())
        cust_options = [f'★ {n}' for n in custom_tpls.keys()]
        all_options  = ['— pick a template —'] + bi_options + cust_options
        tpl_sel = st.selectbox('Template', all_options, key=f'tpl_{sid}',
                               label_visibility='collapsed')

        if tpl_sel != '— pick a template —':
            is_custom   = tpl_sel.startswith('★ ')
            actual_name = tpl_sel[2:] if is_custom else tpl_sel

            if is_custom:
                tpl_data = custom_tpls[actual_name]
                mkts = ', '.join(MARKET_LABELS.get(m, m) for m in tpl_data.get('markets', []))
                st.caption(f'€{tpl_data.get("budget", 0):,.0f} · {mkts}')
                ba, bb = st.columns([3, 1])
                if ba.button('Apply', key=f'tpl_apply_cust_{sid}', use_container_width=True):
                    _apply_template_data(sid, tpl_data)
                safe = actual_name.replace(' ', '_')[:24]
                if bb.button('✕ Delete', key=f'del_tpl_{safe}_{sid}', use_container_width=True):
                    del st.session_state['custom_templates'][actual_name]
                    st.session_state['_pending_dup_rerun'] = True
            else:
                if st.button(f'Apply "{actual_name}"', key=f'tpl_apply_bi_{sid}'):
                    _apply_template(sid, actual_name)

        st.divider()

        # ── Save current as template ──────────────────────────────────────────
        st.markdown('<small style="color:#6b7280;font-weight:600;text-transform:uppercase;'
                    'letter-spacing:0.06em">Save current as template</small>',
                    unsafe_allow_html=True)
        _tpl_key = f'tpl_name_{sid}'
        if st.session_state.pop(f'_tpl_name_clear_{sid}', False):
            st.session_state[_tpl_key] = ''

        new_tpl_name = st.text_input('Template name', key=_tpl_key,
                                     placeholder='e.g. Q3 DACH Full Funnel',
                                     label_visibility='collapsed')
        if st.button('💾 Save template', key=f'save_tpl_{sid}',
                     use_container_width=True, disabled=not new_tpl_name.strip()):
            name = new_tpl_name.strip()
            st.session_state['custom_templates'][name] = _current_as_template(sid)
            st.session_state[f'_tpl_name_clear_{sid}'] = True
            st.session_state['_pending_dup_rerun'] = True

    st.divider()

    # ── Goals & Channels ─────────────────────────────────────────────────────
    st.markdown('**Goals & Channels**')
    hc = st.columns([2, 1, 1, 1, 1])
    hc[1].markdown('<small>YT</small>', unsafe_allow_html=True)
    hc[2].markdown('<small>Display</small>', unsafe_allow_html=True)
    hc[3].markdown('<small>Search</small>', unsafe_allow_html=True)
    hc[4].markdown('<small>LinkedIn</small>', unsafe_allow_html=True)

    goal_channels = {}
    for goal_key in ALL_GOALS:
        rc = st.columns([2, 1, 1, 1, 1])
        goal_on  = rc[0].checkbox(goal_key, value=False, key=f'sb_goal_{goal_key}_{sid}')
        yt_on    = rc[1].checkbox('YT',  value=False, key=f'sb_yt_{goal_key}_{sid}',  label_visibility='collapsed', disabled=not goal_on)
        dis_on   = rc[2].checkbox('Dis', value=False, key=f'sb_dis_{goal_key}_{sid}', label_visibility='collapsed', disabled=not goal_on)
        s_on     = rc[3].checkbox('S',   value=False, key=f'sb_s_{goal_key}_{sid}',   label_visibility='collapsed', disabled=not goal_on)
        li_on    = rc[4].checkbox('LI',  value=False, key=f'sb_li_{goal_key}_{sid}',  label_visibility='collapsed', disabled=not goal_on)
        if goal_on:
            chs = [ch for ch, on in [('YouTube', yt_on), ('Display', dis_on), ('Search', s_on), ('LinkedIn', li_on)] if on]
            if chs:
                goal_channels[goal_key] = chs

    if not goal_channels:
        st.info('Select at least one goal and channel above.')
        return None

    selected_goals = list(goal_channels.keys())
    st.divider()

    # ── Markets & Budget ─────────────────────────────────────────────────────
    cfg1, cfg2 = st.columns([4, 1])
    with cfg1:
        st.markdown('**Markets**')
        # Market group shortcuts
        grp_cols = st.columns(len(MARKET_GROUPS))
        for gi, (grp_label, grp_mkts) in enumerate(MARKET_GROUPS.items()):
            if grp_cols[gi].button(grp_label, key=f'grp_{grp_label}_{sid}',
                                   use_container_width=True, help=f'Add {", ".join(grp_mkts)}'):
                current = list(st.session_state.get(f'selected_markets_{sid}', []))
                merged  = current + [m for m in grp_mkts if m not in current]
                st.session_state[f'selected_markets_{sid}'] = merged
                st.session_state['_pending_dup_rerun'] = True
        s_markets = st.multiselect(
            'Markets', list(MARKET_LABELS.keys()), default=[],
            format_func=lambda k: f'{k} — {MARKET_LABELS[k]}',
            label_visibility='collapsed',
            key=f'selected_markets_{sid}',
        )
    with cfg2:
        st.markdown('**Total Budget (€)**')
        s_budget = st.number_input(
            'Budget', min_value=0, value=0, step=500,
            label_visibility='collapsed', key=f'total_budget_{sid}'
        )

    s_market_budgets, s_market_pcts = {}, {}
    if s_markets:
        n_mkts = len(s_markets)
        default_pct = round(100.0 / n_mkts, 1)

        # Budget split shortcuts
        sc1, sc2, _ = st.columns([1, 1, 3])
        if sc1.button('⚖ Equal split', key=f'eq_{sid}', use_container_width=True):
            for m in s_markets:
                st.session_state[f'pct_{m}_{sid}'] = default_pct
            st.session_state['_pending_dup_rerun'] = True
        if sc2.button('📊 CPM-efficient', key=f'cpm_eff_{sid}', use_container_width=True,
                      help='More budget to cheaper markets to maximise impressions'):
            weights = {}
            for m in s_markets:
                cpms = [BENCH[m][ch]['cpm'] for ch in ['YouTube', 'LinkedIn', 'Display']
                        if ch in BENCH[m] and 'cpm' in BENCH[m][ch]]
                weights[m] = 1.0 / (sum(cpms) / len(cpms)) if cpms else 0.1
            total_w = sum(weights.values()) or 1
            for m in s_markets:
                st.session_state[f'pct_{m}_{sid}'] = round(weights[m] / total_w * 100, 1)
            st.session_state['_pending_dup_rerun'] = True

        # Keep bud_mkt keys in sync when total budget changes or a new market appears
        _last_key = f'_last_total_{sid}'
        if st.session_state.get(_last_key) != s_budget:
            for m in s_markets:
                pct = st.session_state.get(f'pct_{m}_{sid}', default_pct)
                st.session_state[f'bud_mkt_{m}_{sid}'] = round(s_budget * pct / 100)
            st.session_state[_last_key] = s_budget
        else:
            for m in s_markets:
                if f'bud_mkt_{m}_{sid}' not in st.session_state:
                    pct = st.session_state.get(f'pct_{m}_{sid}', default_pct)
                    st.session_state[f'bud_mkt_{m}_{sid}'] = round(s_budget * pct / 100)

        # Split inputs + donut side by side
        split_col, donut_col = st.columns([3, 2])
        with split_col:
            st.markdown('**Market Split**')
            per_row = min(n_mkts, 3)
            for row_start in range(0, n_mkts, per_row):
                row_mkts = s_markets[row_start:row_start + per_row]
                row_cols = st.columns(len(row_mkts) * 2)
                for i, mkt in enumerate(row_mkts):
                    pct = row_cols[i * 2].number_input(
                        f'{mkt} %', min_value=0.0, max_value=100.0,
                        value=default_pct, step=0.5, format='%.1f',
                        key=f'pct_{mkt}_{sid}',
                        on_change=_sync_pct_to_bud, args=(mkt, sid),
                    )
                    s_market_pcts[mkt] = pct
                    s_market_budgets[mkt] = s_budget * pct / 100
                    row_cols[i * 2 + 1].number_input(
                        f'{mkt} €', min_value=0, step=100, format='%d',
                        value=st.session_state.get(f'bud_mkt_{mkt}_{sid}', 0),
                        key=f'bud_mkt_{mkt}_{sid}',
                        on_change=_sync_bud_to_pct, args=(mkt, sid),
                    )
            pct_sum = sum(s_market_pcts.values())
            if abs(pct_sum - 100) > 0.5:
                st.warning(f'Split: {pct_sum:.1f}% — adjust to 100%')
            else:
                st.success(f'✓ €{s_budget:,} allocated')
        with donut_col:
            if s_market_pcts:
                st.plotly_chart(_market_donut(s_market_pcts), use_container_width=True,
                                config={'displayModeBar': False}, key=f'donut_{sid}')
    else:
        st.info('Select at least one market above to build the plan.')
        return None

    # Pacing chart (inline, compact)
    _pacing_chart(periods, s_budget, sid)

    st.divider()

    # ── Pinned country panel ──────────────────────────────────────────────────
    pinned_mkt = st.session_state.get(f'pinned_country_{sid}')
    if pinned_mkt and pinned_mkt in s_markets:
        pm_budget = s_market_budgets.get(pinned_mkt, 0)
        pm_pct    = s_market_pcts.get(pinned_mkt, 0)
        pm_goals  = '  ·  '.join(
            f'{g}: {", ".join(chs)}' for g, chs in goal_channels.items()
        )
        cached    = st.session_state.get(f'cached_kpis_{pinned_mkt}_{sid}', {})
        kpi_html  = ('  <span style="color:#666">|</span>  '.join(
            f'<b>{k}</b>: {v}' for k, v in cached.items()
        )) if cached else '<span style="color:#888;font-style:italic">KPIs appear after first render</span>'
        st.markdown(
            '<style>'
            '.pinned-bar{position:sticky;top:3.6rem;z-index:200;background:#f0faf9;'
            'border-left:4px solid #2BB5A5;border-radius:0 6px 6px 0;'
            'padding:7px 14px;margin-bottom:10px;box-shadow:0 2px 8px rgba(0,0,0,0.08);'
            'font-size:0.82rem;line-height:1.5}'
            '</style>'
            f'<div class="pinned-bar">📌 <strong>{MARKET_LABELS[pinned_mkt]}</strong>'
            f'&nbsp;—&nbsp;€{pm_budget:,.0f} ({pm_pct:.0f}%)'
            f'&nbsp;&nbsp;<span style="color:#2BB5A5;font-size:0.75rem">{pm_goals}</span>'
            f'<br>{kpi_html}</div>',
            unsafe_allow_html=True,
        )

    # ── Plan tables ───────────────────────────────────────────────────────────
    grand_totals = {g: {ch: [] for ch in chs} for g, chs in goal_channels.items()}

    for mkt in s_markets:
        mkt_budget = s_market_budgets[mkt]
        is_pinned  = (st.session_state.get(f'pinned_country_{sid}') == mkt)
        mkt_label  = MARKET_LABELS[mkt]

        # Expanded state stored in session state so reruns (pin/unpin) don't reset it
        exp_key = f'mkt_exp_{mkt}_{sid}'
        if exp_key not in st.session_state:
            st.session_state[exp_key] = True
        is_expanded = st.session_state[exp_key]

        # Country header row: toggle + name + pin indicator
        h_col, tog_col, pin_col = st.columns([10, 1, 1])
        pin_suffix = '  📌' if is_pinned else ''
        h_col.subheader(f'{mkt_label}{pin_suffix}')
        tog_icon = '▼' if is_expanded else '▶'
        if tog_col.button(tog_icon, key=f'mkt_tog_{mkt}_{sid}', use_container_width=True,
                          help='Collapse' if is_expanded else 'Expand'):
            st.session_state[exp_key] = not is_expanded
            st.session_state['_pending_dup_rerun'] = True
        pin_icon = '✕ Unpin' if is_pinned else '📌 Pin'
        pin_tip  = 'Unpin this country' if is_pinned else 'Pin to top for comparison while scrolling'
        if pin_col.button(pin_icon, key=f'pin_{mkt}_{sid}',
                          use_container_width=True, help=pin_tip):
            if is_pinned:
                st.session_state.pop(f'pinned_country_{sid}', None)
            else:
                st.session_state[f'pinned_country_{sid}'] = mkt
            st.session_state['_pending_dup_rerun'] = True

        if is_expanded:
            # Goal budget split — only shown when multiple goals are selected
            goal_budgets = {}
            if len(selected_goals) > 1:
                default_goal_pct = round(100.0 / len(selected_goals), 1)
                st.markdown('**Goal budget split**')
                gcols = st.columns(len(selected_goals))
                raw_pcts = {}
                for gi, goal in enumerate(selected_goals):
                    raw_pcts[goal] = gcols[gi].number_input(
                        f'{goal} %', min_value=0.0, max_value=100.0,
                        value=default_goal_pct, step=5.0, format='%.1f',
                        key=f'goal_pct_{mkt}_{goal}_{sid}',
                    )
                pct_sum_g = sum(raw_pcts.values()) or 1
                for gi, goal in enumerate(selected_goals):
                    goal_budgets[goal] = mkt_budget * raw_pcts[goal] / pct_sum_g
                    gcols[gi].markdown(
                        f'<div style="font-size:0.85rem;font-weight:600;color:#2BB5A5">'
                        f'€{goal_budgets[goal]:,.0f}</div>',
                        unsafe_allow_html=True,
                    )
                if abs(sum(raw_pcts.values()) - 100) > 0.5:
                    st.caption(f'Goal split sums to {sum(raw_pcts.values()):.1f}% — normalising proportionally.')
            else:
                goal_budgets[selected_goals[0]] = mkt_budget

            if len(selected_goals) > 1:
                goal_tabs = st.tabs(selected_goals)
                for tab, goal in zip(goal_tabs, selected_goals):
                    with tab:
                        goal_chs = goal_channels[goal]
                        ch_budgets = _channel_budget_split(mkt, goal, goal_chs, goal_budgets[goal], sid)
                        render_goal_section(mkt, goal, goal_chs, ch_budgets, periods, grand_totals, sid)
            else:
                goal = selected_goals[0]
                goal_chs = goal_channels[goal]
                ch_budgets = _channel_budget_split(mkt, goal, goal_chs, goal_budgets[goal], sid)
                render_goal_section(mkt, goal, goal_chs, ch_budgets, periods, grand_totals, sid)

            # Cache KPI totals for the pinned panel (available on next render)
            if mkt == st.session_state.get(f'pinned_country_{sid}'):
                kpi_cache = {}
                for g in selected_goals:
                    for ch in goal_channels[g]:
                        rows = grand_totals[g][ch]
                        if rows:
                            t = rows[-1].iloc[0]
                            parts = [
                                f'{COL_FMT[c][0]}: {COL_FMT[c][1](t[c])}'
                                for c in ['impressions', 'reach', 'views', 'clicks', 'sessions', 'conversions']
                                if c in t.index and t[c] > 0
                            ]
                            if parts:
                                kpi_cache[f'{g} · {ch}'] = ' | '.join(parts[:3])
                st.session_state[f'cached_kpis_{pinned_mkt}_{sid}'] = kpi_cache

        st.divider()

    st.subheader('Grand Total — All Markets')
    for goal in selected_goals:
        if len(selected_goals) > 1:
            st.markdown(f'### {goal}')
        for ch in goal_channels[goal]:
            rows = grand_totals[goal][ch]
            if not rows:
                continue
            gdf = pd.concat(rows, ignore_index=True)
            st.markdown(f'**{ch} — per market**')
            disp = gdf[['Market']].copy()
            for col in ADDITIVE:
                if col in gdf.columns and col in COL_FMT:
                    label, fn = COL_FMT[col]
                    disp[label] = gdf[col].apply(fn)
            st.dataframe(disp, use_container_width=True, hide_index=True)
            grand = {'Period': 'TOTAL', 'Days': int(gdf['Days'].sum())}
            for col in ADDITIVE:
                if col in gdf.columns:
                    grand[col] = gdf[col].sum()
            grand_df = pd.DataFrame([grand])
            g1, g2 = st.columns([2, 3])
            with g1:
                st.markdown(f'**{ch} — combined totals**')
                rows_disp = [{'Metric': COL_FMT[c][0], 'Value': COL_FMT[c][1](grand[c])}
                             for c in ADDITIVE if c in grand and c in COL_FMT]
                st.dataframe(pd.DataFrame(rows_disp), use_container_width=True, hide_index=True)
            with g2:
                st.plotly_chart(
                    make_funnel(grand_df, goal, ch, f'Grand Total — {ch} ({goal})'),
                    use_container_width=True, config={'displayModeBar': False},
                    key=f'grand_{goal}_{ch}_{sid}',
                )
        if len(selected_goals) > 1:
            st.divider()

    return {
        'name': st.session_state.scenario_names[sid],
        'sid': sid,
        'grand_totals': grand_totals,
        'goal_channels': goal_channels,
        'selected_markets': s_markets,
        'market_budgets': s_market_budgets,
        'market_pcts': s_market_pcts,
        'grand_total_bud': s_budget,
    }


_scenario_sids = []
for sid, s_tab in enumerate(scenario_tabs):
    with s_tab:
        s_data = _render_scenario(sid)
        if s_data:
            all_scenarios_data.append(s_data)
            _scenario_sids.append(sid)
            st.divider()
            gads_bytes = _build_gads_csv_scenario(s_data, sid, campaign_name, start_date, end_date)
            st.download_button(
                label=f'⬇ Google Ads CSV — {s_data["name"]}',
                data=gads_bytes,
                file_name=f'{campaign_name.replace(" ", "_")}_{s_data["name"].replace(" ", "_")}_google_ads.csv',
                mime='text/csv',
                key=f'dl_gads_{sid}',
                use_container_width=True,
            )

# Deferred rerun from _duplicate_scenario — all tabs have now fully rendered,
# so no widget state will be orphaned by this rerun.
if st.session_state.pop('_pending_dup_rerun', False):
    st.rerun()

# ── Combined Excel download (all scenarios as tabs) ───────────────────────────
if all_scenarios_data:
    xl_bytes = _build_excel_all(all_scenarios_data, _scenario_sids, campaign_name, start_date, end_date, breakdown)
    st.download_button(
        label=f'⬇ Download Excel — All Scenarios ({len(all_scenarios_data)} tab{"s" if len(all_scenarios_data) != 1 else ""})',
        data=xl_bytes,
        file_name=f'{campaign_name.replace(" ", "_")}_media_plan.xlsx',
        mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        key='dl_excel_all',
        use_container_width=True,
    )


# ── Compare tab (only when 2+ scenarios) ─────────────────────────────────────

def _aggregate_scenario_metrics(s):
    """Sum all additive metrics across goals and channels for a scenario."""
    totals = {col: 0 for col in ADDITIVE}
    for goal, chs in s['grand_totals'].items():
        for ch, rows in chs.items():
            if not rows:
                continue
            gdf = pd.concat(rows, ignore_index=True)
            for col in ADDITIVE:
                if col in gdf.columns:
                    totals[col] += gdf[col].sum()
    return totals


def _kpi_rows(s, metrics):
    """Build text lines per goal+channel for the AI prompt."""
    lines = []
    for goal, chs in s['grand_totals'].items():
        for ch, rows in chs.items():
            if not rows:
                continue
            gdf = pd.concat(rows, ignore_index=True)
            parts = []
            for col in ADDITIVE:
                if col in gdf.columns and gdf[col].sum() > 0:
                    parts.append(f'{COL_FMT[col][0]}: {COL_FMT[col][1](gdf[col].sum())}')
            if parts:
                lines.append(f'    [{goal} / {ch}] {" | ".join(parts)}')
    return '\n'.join(lines)


if compare_tab is not None:
    with compare_tab:
        st.caption('Side-by-side KPI comparison and AI recommendation — budget efficiency, reach, and funnel performance.')
        if len(all_scenarios_data) < 2:
            st.info('Fill in at least two scenarios (markets, goals, channels, and budgets) to run a comparison.')
        else:
            # ── Scenario selector ─────────────────────────────────────────────
            all_names = [s['name'] for s in all_scenarios_data]
            selected_names = st.multiselect(
                'Scenarios to compare',
                options=all_names,
                default=all_names,
                key='compare_selected',
            )
            compare_data = [s for s in all_scenarios_data if s['name'] in selected_names]

            if len(compare_data) < 2:
                st.info('Select at least two scenarios to compare.')
            else:
                metric_labels = {col: COL_FMT[col][0] for col in ADDITIVE if col in COL_FMT}

                # ── KPI summary table ─────────────────────────────────────────
                st.markdown('#### KPI Summary')
                summary_rows = []
                for s in compare_data:
                    agg = _aggregate_scenario_metrics(s)
                    row = {'Scenario': s['name'], 'Budget (€)': f"€{agg['Budget']:,.0f}"}
                    for col in ADDITIVE:
                        if col == 'Budget':
                            continue
                        if agg.get(col, 0) > 0:
                            row[metric_labels[col]] = COL_FMT[col][1](agg[col])
                    summary_rows.append(row)
                st.dataframe(pd.DataFrame(summary_rows), use_container_width=True, hide_index=True)

                # ── Winner callouts per metric ────────────────────────────────
                st.markdown('#### Which scenario leads on each KPI?')
                active_cols = [c for c in ADDITIVE if c != 'Budget'
                               and any(_aggregate_scenario_metrics(s).get(c, 0) > 0 for s in compare_data)]
                winner_cols = st.columns(max(len(active_cols), 1))
                for col_idx, col in enumerate(active_cols):
                    vals = [(s['name'], _aggregate_scenario_metrics(s).get(col, 0)) for s in compare_data]
                    best = max(vals, key=lambda x: x[1])
                    winner_cols[col_idx].markdown(
                        f'<div style="border:1px solid #e2e8f0;border-radius:8px;padding:10px 12px;'
                        f'text-align:center;background:#f8fafc;">'
                        f'<div style="font-size:0.70rem;color:#64748b;font-weight:600;'
                        f'text-transform:uppercase;letter-spacing:0.05em;margin-bottom:4px">'
                        f'{metric_labels[col]}</div>'
                        f'<div style="font-size:0.82rem;font-weight:700;color:#1e293b;'
                        f'line-height:1.3;word-break:break-word">{best[0]}</div>'
                        f'</div>',
                        unsafe_allow_html=True,
                    )

                # ── Visual comparison bar charts ──────────────────────────────
                st.markdown('#### Visual comparison')
                bar_metrics = [c for c in ['impressions', 'reach', 'clicks', 'sessions', 'conversions']
                               if any(_aggregate_scenario_metrics(s).get(c, 0) > 0 for s in compare_data)]
                if bar_metrics:
                    sc_names  = [s['name'] for s in compare_data]
                    sc_colors = DONUT_PALETTE[:len(compare_data)]
                    sc_labels = [n if len(n) <= 18 else n[:16] + '…' for n in sc_names]
                    bar_cols  = st.columns(min(len(bar_metrics), 3))
                    for bi, metric in enumerate(bar_metrics[:6]):
                        with bar_cols[bi % 3]:
                            vals = [_aggregate_scenario_metrics(s).get(metric, 0) for s in compare_data]
                            fig  = go.Figure(go.Bar(
                                x=sc_labels, y=vals,
                                customdata=sc_names,
                                hovertemplate='<b>%{customdata}</b><br>' + COL_FMT[metric][0] + ': %{text}<extra></extra>',
                                marker_color=sc_colors,
                                text=[COL_FMT[metric][1](v) for v in vals],
                                textposition='outside',
                                textfont={'size': 11},
                            ))
                            fig.update_layout(
                                title={'text': COL_FMT[metric][0],
                                       'font': {'size': 14, 'color': '#1e293b'}, 'x': 0.5, 'xanchor': 'center'},
                                margin={'t': 40, 'b': 60, 'l': 10, 'r': 10},
                                height=230,
                                paper_bgcolor='rgba(0,0,0,0)',
                                yaxis={'showticklabels': False, 'showgrid': False, 'zeroline': False},
                                xaxis={'tickfont': {'size': 12}, 'tickangle': -20},
                            )
                            st.plotly_chart(fig, use_container_width=True,
                                            config={'displayModeBar': False}, key=f'bar_{metric}')

                st.divider()

                # ── AI deep-dive ──────────────────────────────────────────────
                if st.button('Generate AI Comparison', key='btn_compare'):
                    api_key = get_api_key()
                    if not api_key:
                        st.error('No API key found in .streamlit/secrets.toml')
                    else:
                        def _scenario_block(s):
                            agg = _aggregate_scenario_metrics(s)
                            lines = [f"Scenario: {s['name']}"]
                            lines.append(f"  Total budget: €{s['grand_total_bud']:,.0f}")
                            lines.append(f"  Markets ({len(s['selected_markets'])}): {', '.join(MARKET_LABELS[m] for m in s['selected_markets'])}")
                            for goal, chs in s['goal_channels'].items():
                                lines.append(f"  Goal: {goal} | Channels: {', '.join(chs)}")
                            lines.append('  Aggregated KPIs:')
                            for col in ADDITIVE:
                                if col in agg and agg[col] > 0:
                                    lines.append(f'    {COL_FMT[col][0]}: {COL_FMT[col][1](agg[col])}')
                            lines.append('  KPIs by goal and channel:')
                            lines.append(_kpi_rows(s, agg))
                            return '\n'.join(lines)

                        scenarios_text = '\n\n'.join(_scenario_block(s) for s in compare_data)
                        compare_prompt = f"""You are a senior paid media strategist comparing {len(compare_data)} media plan scenarios for a {audience_type} brand in the {industry} sector.

{scenarios_text}

For each scenario write a clearly labelled block using exactly this format (one block per scenario):

SCENARIO: [name]
STRENGTHS: [60–80 words — what this scenario does well for a {audience_type} {industry} brand: market coverage, channel fit, which KPIs it leads on and why that matters for this industry]
WEAKNESSES: [60–80 words — where this scenario underperforms: KPIs it loses on, missing markets, channel imbalance, or poor fit with {audience_type} buyer journeys in {industry}]
VERDICT: [25–35 words — one direct sentence on the single best use case for this scenario]

After all scenario blocks, write these two final sections:

BEST FOR BUDGET EFFICIENCY:
[50–60 words — which scenario delivers the most value per euro spent, referencing the actual KPI numbers. Frame this for a {audience_type} {industry} brand where budget scrutiny is typical.]

BEST FOR KPI PERFORMANCE:
[50–60 words — which scenario wins on raw KPI delivery per funnel stage (awareness → traffic → conversion), and what that means for a {audience_type} {industry} campaign objective.]

No bullet points. Write like a strategist presenting to a client."""

                        import anthropic as _anthropic
                        client = _anthropic.Anthropic(api_key=api_key)
                        with st.spinner('Comparing scenarios...'):
                            msg = client.messages.create(
                                model='claude-sonnet-4-6',
                                max_tokens=1100,
                                messages=[{'role': 'user', 'content': compare_prompt}]
                            )
                        raw = msg.content[0].text
                        import re as _re
                        blocks = _re.split(r'\n(SCENARIO|STRENGTHS|WEAKNESSES|VERDICT|BEST FOR BUDGET EFFICIENCY|BEST FOR KPI PERFORMANCE):\s*', raw)
                        if len(blocks) > 1:
                            labels = blocks[1::2]
                            bodies = blocks[2::2]
                            palette = {
                                'SCENARIO':                   ('#1F497D', 'white'),
                                'STRENGTHS':                  ('#1F6152', 'white'),
                                'WEAKNESSES':                 ('#8B3A3A', 'white'),
                                'VERDICT':                    ('#437CA3', 'white'),
                                'BEST FOR BUDGET EFFICIENCY': ('#5A3E7A', 'white'),
                                'BEST FOR KPI PERFORMANCE':   ('#2D6A8A', 'white'),
                            }
                            for label, body in zip(labels, bodies):
                                bg, fg = palette.get(label, ('#555', 'white'))
                                st.markdown(
                                    f'<div style="background:{bg};color:{fg};padding:6px 12px;border-radius:4px 4px 0 0;'
                                    f'font-weight:bold;margin-top:10px">{label}</div>'
                                    f'<div style="background:#f5f5f5;padding:10px 12px;border-radius:0 0 4px 4px;'
                                    f'margin-bottom:2px">{body.strip()}</div>',
                                    unsafe_allow_html=True,
                                )
                        else:
                            st.markdown(raw)


# ── AI Insights & Recommendations ────────────────────────────────────────────
st.divider()
st.subheader('AI Insights & Recommendations')

if all_scenarios_data:
    scenario_options = [s['name'] for s in all_scenarios_data]
    ai_scenario_name = st.selectbox('Scenario to analyse', scenario_options)
    ai_data = next(s for s in all_scenarios_data if s['name'] == ai_scenario_name)
else:
    ai_data = None

if ai_data:
    _ai_channels = sorted({ch for chs in ai_data['goal_channels'].values() for ch in chs})
    _ai_markets  = [MARKET_LABELS[m] for m in ai_data['selected_markets']]
    st.caption(
        f'Insights framed for **{audience_type}** · **{industry}** · '
        f'{", ".join(_ai_channels)} · {", ".join(_ai_markets)}'
    )
else:
    st.caption(f'Insights framed for **{audience_type}** · **{industry}**')

def build_plan_summary():
    """Compile a plain-text plan summary from the selected scenario's data, using
    the actual user-edited benchmark values from session state."""
    if not ai_data:
        return 'No scenario data available.'
    d   = ai_data
    sid = d.get('sid', 0)
    _is_pct = {'ctr', 'view_rate', 'click_to_session', 'conv_rate', 'lead_to_mql', 'mql_to_sql'}

    def _ss_bench(key, mkt, ch, goal):
        """Read a benchmark value from session state; fall back to BENCH default."""
        raw = st.session_state.get(f'{key}_{mkt}_{ch}_{goal}_{sid}')
        if raw is not None:
            return raw / 100.0 if key in _is_pct else raw
        return BENCH[mkt][ch].get(key)

    lines = [
        f'Campaign: {campaign_name}  ({d["name"]})',
        f'Audience: {audience_type}  |  Industry: {industry}',
        f'Flight: {start_date.strftime("%b %d, %Y")} – {end_date.strftime("%b %d, %Y")} ({(end_date - start_date).days + 1} days)',
        f'Total budget: €{d["grand_total_bud"]:,}',
        '',
        'Market budgets:',
    ]
    n_goals = len(d['goal_channels'])
    for mkt in d['selected_markets']:
        mkt_bud = d['market_budgets'].get(mkt, 0)
        mkt_pct = d['market_pcts'].get(mkt, 0)
        lines.append(f'  {MARKET_LABELS[mkt]}: €{mkt_bud:,.0f} ({mkt_pct:.1f}%)')
        if n_goals > 1:
            for goal in d['goal_channels']:
                gpct = st.session_state.get(f'goal_pct_{mkt}_{goal}_{sid}',
                                            round(100.0 / n_goals, 1))
                gbud = mkt_bud * gpct / 100
                lines.append(f'    → {goal}: €{gbud:,.0f} ({gpct:.0f}%)')
    lines.append('')

    for goal, chs in d['goal_channels'].items():
        lines.append(f'Goal: {goal}  |  Channels: {", ".join(chs)}')
        for ch in chs:
            # Actual KPI totals
            rows = d['grand_totals'][goal][ch]
            if rows:
                gdf  = pd.concat(rows, ignore_index=True)
                grand = {col: gdf[col].sum() for col in ADDITIVE if col in gdf.columns}
                parts = [f'{COL_FMT[col][0]}: {COL_FMT[col][1](grand[col])}'
                         for col in ADDITIVE if col in grand and col in COL_FMT]
                lines.append(f'  {ch} KPIs: {" | ".join(parts)}')
            # Actual benchmarks used (from session state, not hardcoded defaults)
            bm_parts = []
            if ch == 'Search':
                cpc = _ss_bench('cpc', d['selected_markets'][0], ch, goal)
                ctr = _ss_bench('ctr', d['selected_markets'][0], ch, goal)
                if cpc: bm_parts.append(f'CPC €{cpc:.2f}')
                if ctr: bm_parts.append(f'CTR {ctr*100:.1f}%')
            else:
                cpm  = _ss_bench('cpm', d['selected_markets'][0], ch, goal)
                ctr  = _ss_bench('ctr', d['selected_markets'][0], ch, goal)
                freq = _ss_bench('frequency', d['selected_markets'][0], ch, goal)
                if cpm:  bm_parts.append(f'CPM €{cpm:.2f}')
                if ctr:  bm_parts.append(f'CTR {ctr*100:.2f}%')
                if freq: bm_parts.append(f'Frequency {freq:.1f}')
                if ch == 'YouTube':
                    vr = _ss_bench('view_rate', d['selected_markets'][0], ch, goal)
                    if vr: bm_parts.append(f'View Rate {vr*100:.0f}%')
            if bm_parts:
                lines.append(f'  {ch} benchmarks (first market as reference): {", ".join(bm_parts)}')
        lines.append('')
    return '\n'.join(lines)

# Curated benchmark source list — only cite from this list, never invent URLs
_BENCHMARK_SOURCES = """
YOUTUBE
- WordStream Google Ads Benchmarks 2025: https://www.wordstream.com/blog/2025-google-ads-benchmarks
- YouTube CPM Europe by Country 2026: https://fluxnote.io/guides/youtube-cpm-europe-by-country-2026
- YouTube Ad Benchmarks (CPV, CPM, CTR): https://megadigital.ai/en/blog/youtube-ad-benchmarks/
- YouTube Ads Benchmarks by Industry: https://www.storegrowers.com/youtube-ads-benchmarks/
- YouTube CPM by Country Global Comparison 2026: https://upgrowth.in/youtube-cpm-by-country-global-comparison-2026/
- YouTube CPM by Niche 2026: https://upgrowth.in/youtube-cpm-overview-highest-paying-niches-2026/

LINKEDIN
- LinkedIn Marketing Solutions Blog (official): https://business.linkedin.com/marketing-solutions/blog
- Dreamdata LinkedIn Ads B2B Benchmarks: https://dreamdata.io/linkedin-ads-b2b-benchmarks

SEARCH & DISPLAY
- WordStream Google Ads Benchmarks 2025: https://www.wordstream.com/blog/2025-google-ads-benchmarks
- HubSpot Marketing Statistics 2026: https://www.hubspot.com/marketing-statistics

B2B & GENERAL
- B2B Paid Awareness Benchmarks Q1 2026: https://refinelabs.com/article/b2b-paid-awareness-benchmarks-q1-2026
- Paid Social Advertising Costs UK 2026: https://www.mediaperformance.co.uk/paid-social-advertising-costs-uk-2026/
- Hootsuite Social Media Benchmarks 2026: https://blog.hootsuite.com/social-media-benchmarks/
"""

ai_tab_insights, ai_tab_recs, ai_tab_benchmarks = st.tabs([
    'Plan Insights', 'Market Recommendations', 'Benchmark Explanations'
])

with ai_tab_insights:
    st.caption('High-level read of the plan — what the numbers say, what looks strong, what to watch.')
    if st.button('Generate Plan Insights', key='btn_insights'):
        api_key = get_api_key()
        if not api_key:
            st.error('No API key found in .streamlit/secrets.toml')
        else:
            summary = build_plan_summary()
            prompt = f"""You are a senior paid media strategist reviewing a digital media plan for a {audience_type} brand in the {industry} sector.

{summary}

Write a concise but sharp analysis (150–200 words). Cover:
1. Overall scale and reach potential given the budget and markets — frame this for {audience_type} {industry} audiences specifically
2. Channel mix strengths or gaps for the selected goals — comment on what works for {audience_type} in this sector
3. Any markets where the allocation looks strong or under-invested for this industry
4. One clear watch-out or risk relevant to {audience_type} {industry} campaigns

Be direct. No headers. No bullet points. Write in plain paragraphs like a strategist talking to a client.

After the analysis, add a 'Sources' section listing 2–4 of the most relevant URLs from the reference list below that support your specific points. Copy the URLs exactly as written — do not modify or invent any URL.

Reference sources:
{_BENCHMARK_SOURCES}"""
            import anthropic as _anthropic
            client = _anthropic.Anthropic(api_key=api_key)
            with st.spinner('Thinking...'):
                msg = client.messages.create(
                    model='claude-sonnet-4-6',
                    max_tokens=600,
                    messages=[{'role': 'user', 'content': prompt}]
                )
            st.session_state['insights_last'] = msg.content[0].text
        st.rerun()

    if st.session_state.get('insights_last'):
        st.markdown(st.session_state['insights_last'])

    st.divider()
    st.markdown('#### Chat about the plan')
    st.caption('Ask anything — follow up, dig deeper, change direction. The conversation keeps its memory.')

    if 'insights_chat' not in st.session_state:
        st.session_state.insights_chat = []

    if st.session_state.insights_chat:
        if st.button('🗑 Clear conversation', key='insights_clear'):
            st.session_state.insights_chat = []
            st.rerun()

    for msg in st.session_state.insights_chat:
        with st.chat_message(msg['role']):
            st.markdown(msg['content'])

    if user_input := st.chat_input(
        'e.g. What if we cut the Netherlands budget by 20%? Which channel is working hardest?',
        key='insights_chat_input',
    ):
        with st.chat_message('user'):
            st.markdown(user_input)

        api_key = get_api_key()
        if not api_key:
            with st.chat_message('assistant'):
                st.error('No API key found in .streamlit/secrets.toml')
        else:
            summary = build_plan_summary()
            system_ctx = (
                f'You are a senior paid media strategist advising a {audience_type} client '
                f'in the {industry} sector. Here is the full media plan:\n\n{summary}\n\n'
                f'Answer in 80–150 words. Be direct, reference actual numbers from the plan, '
                f'frame everything for {audience_type} {industry}. No bullet points.'
            )
            api_messages = []
            for i, m in enumerate(st.session_state.insights_chat):
                if i == 0 and m['role'] == 'user':
                    api_messages.append({'role': 'user', 'content': system_ctx + '\n\n---\n\n' + m['content']})
                else:
                    api_messages.append(m)
            first_turn = len(st.session_state.insights_chat) == 0
            api_messages.append({
                'role': 'user',
                'content': (system_ctx + '\n\n---\n\n' + user_input) if first_turn else user_input,
            })
            import anthropic as _anthropic
            client = _anthropic.Anthropic(api_key=api_key)
            with st.chat_message('assistant'):
                with st.spinner(''):
                    resp = client.messages.create(
                        model='claude-sonnet-4-6',
                        max_tokens=400,
                        messages=api_messages,
                    )
                reply = resp.content[0].text
                st.markdown(reply)

            st.session_state.insights_chat.append({'role': 'user',      'content': user_input})
            st.session_state.insights_chat.append({'role': 'assistant', 'content': reply})

with ai_tab_recs:
    st.caption('Budget allocation recommendations — which markets or channels deserve more weight and why.')
    if st.button('Generate Market Recommendations', key='btn_recs'):
        api_key = get_api_key()
        if not api_key:
            st.error('No API key found in .streamlit/secrets.toml')
        else:
            summary = build_plan_summary()
            mkt_list = '\n'.join(
                f'- {MARKET_LABELS[m]}: {ai_data["market_pcts"][m]:.1f}% (€{ai_data["market_budgets"][m]:,.0f})'
                for m in ai_data['selected_markets']
            ) if ai_data else 'No data.'
            channels_in_use = sorted({ch for chs in (ai_data['goal_channels'].values() if ai_data else []) for ch in chs})
            prompt = f"""You are a senior paid media strategist advising a {audience_type} brand in the {industry} sector on market budget allocation across {', '.join(channels_in_use) if channels_in_use else 'digital channels'}.

{summary}

Current market split:
{mkt_list}

Write exactly 3 recommendation sections, each starting with its label on its own line:

CURRENT ALLOCATION:
[50–60 words assessing whether the current market split makes sense for a {audience_type} {industry} brand — consider CPM efficiency vs audience quality for this sector and whether the channels in use ({', '.join(channels_in_use) if channels_in_use else 'the selected channels'}) justify the current weights]

REBALANCING OPTION:
[50–60 words suggesting one concrete rebalancing move — e.g. shift 10% from X to Y — grounded in what works for {audience_type} audiences in {industry} and the channel mix being used]

BEST ALLOCATION:
[50–60 words giving a direct final recommendation on the optimal split, framed specifically for a {audience_type} {industry} brand running {', '.join(channels_in_use) if channels_in_use else 'these channels'}]

SOURCES:
[List 2–3 of the most relevant URLs from the reference list below, copied exactly — do not modify or invent any URL]

Be direct. No bullet points within sections.

Reference sources:
{_BENCHMARK_SOURCES}"""
            import anthropic as _anthropic
            client = _anthropic.Anthropic(api_key=api_key)
            with st.spinner('Thinking...'):
                msg = client.messages.create(
                    model='claude-sonnet-4-6',
                    max_tokens=700,
                    messages=[{'role': 'user', 'content': prompt}]
                )
            st.session_state['recs_last'] = msg.content[0].text
        st.rerun()

    if st.session_state.get('recs_last'):
        raw = st.session_state['recs_last']
        import re as _re
        sections = _re.split(r'\n(CURRENT ALLOCATION|REBALANCING OPTION|BEST ALLOCATION|SOURCES):\n', raw)
        if len(sections) > 1:
            labels = sections[1::2]
            bodies = sections[2::2]
            colors = {
                'CURRENT ALLOCATION': '#437CA3',
                'REBALANCING OPTION': '#437CA3',
                'BEST ALLOCATION':    '#1F6152',
                'SOURCES':            '#555555',
            }
            for label, body in zip(labels, bodies):
                color = colors.get(label, '#437CA3')
                st.markdown(f'<div style="background:{color};color:white;padding:6px 10px;border-radius:4px;font-weight:bold;margin-top:8px">{label}</div>', unsafe_allow_html=True)
                if label == 'SOURCES':
                    st.markdown(body.strip())
                else:
                    st.markdown(f'<div style="background:#f2f2f2;padding:10px;border-radius:0 0 4px 4px;margin-bottom:4px">{body.strip()}</div>', unsafe_allow_html=True)
        else:
            st.markdown(raw)

    st.divider()
    st.markdown('#### Chat about market allocation')
    st.caption('Ask anything — follow up, dig deeper, change direction. The conversation keeps its memory.')

    if 'recs_chat' not in st.session_state:
        st.session_state.recs_chat = []

    if st.session_state.recs_chat:
        if st.button('🗑 Clear conversation', key='recs_clear'):
            st.session_state.recs_chat = []
            st.rerun()

    for msg in st.session_state.recs_chat:
        with st.chat_message(msg['role']):
            st.markdown(msg['content'])

    if user_input := st.chat_input(
        'e.g. Should we shift budget from Germany to France? What if we added LinkedIn to the mix?',
        key='recs_chat_input',
    ):
        with st.chat_message('user'):
            st.markdown(user_input)

        api_key = get_api_key()
        if not api_key:
            with st.chat_message('assistant'):
                st.error('No API key found in .streamlit/secrets.toml')
        else:
            summary = build_plan_summary()
            mkt_list = '\n'.join(
                f'- {MARKET_LABELS[m]}: {ai_data["market_pcts"][m]:.1f}% (€{ai_data["market_budgets"][m]:,.0f})'
                for m in ai_data['selected_markets']
            ) if ai_data else ''
            channels_in_use = sorted({ch for chs in (ai_data['goal_channels'].values() if ai_data else []) for ch in chs})
            system_ctx = (
                f'You are a senior paid media strategist advising a {audience_type} client '
                f'in the {industry} sector. Here is the full media plan:\n\n{summary}\n\n'
                f'Current market split:\n{mkt_list}\n\n'
                f'Channels in use: {", ".join(channels_in_use) if channels_in_use else "various"}. '
                f'Answer in 80–150 words. Be direct, reference actual numbers, '
                f'frame for {audience_type} {industry}. No bullet points.'
            )
            api_messages = []
            for i, m in enumerate(st.session_state.recs_chat):
                if i == 0 and m['role'] == 'user':
                    api_messages.append({'role': 'user', 'content': system_ctx + '\n\n---\n\n' + m['content']})
                else:
                    api_messages.append(m)
            first_turn = len(st.session_state.recs_chat) == 0
            api_messages.append({
                'role': 'user',
                'content': (system_ctx + '\n\n---\n\n' + user_input) if first_turn else user_input,
            })
            import anthropic as _anthropic
            client = _anthropic.Anthropic(api_key=api_key)
            with st.chat_message('assistant'):
                with st.spinner(''):
                    resp = client.messages.create(
                        model='claude-sonnet-4-6',
                        max_tokens=400,
                        messages=api_messages,
                    )
                reply = resp.content[0].text
                st.markdown(reply)

            st.session_state.recs_chat.append({'role': 'user',      'content': user_input})
            st.session_state.recs_chat.append({'role': 'assistant', 'content': reply})

with ai_tab_benchmarks:
    st.caption('Plain-English explanations of the benchmark values used in this plan.')

    def _bench_context():
        """Build a reusable benchmark context block for AI prompts.
        Reads from session state (user-edited values) with BENCH as fallback."""
        if not ai_data:
            return [], []
        sid = ai_data.get('sid', 0)
        _is_pct = {'ctr', 'view_rate', 'click_to_session', 'conv_rate', 'lead_to_mql', 'mql_to_sql'}

        def _get(key, mkt, ch, goal):
            raw = st.session_state.get(f'{key}_{mkt}_{ch}_{goal}_{sid}')
            if raw is not None:
                return raw / 100.0 if key in _is_pct else raw
            return BENCH[mkt][ch].get(key)

        bench_lines = []
        for mkt in ai_data['selected_markets']:
            for goal, chs in ai_data['goal_channels'].items():
                for ch in chs:
                    b = BENCH[mkt][ch]
                    if ch == 'Search':
                        cpc = _get('cpc', mkt, ch, goal) or b.get('cpc', 0)
                        ctr = _get('ctr', mkt, ch, goal) or b.get('ctr', 0)
                        c2s = _get('click_to_session', mkt, ch, goal) or b.get('click_to_session', 0.85)
                        bench_lines.append(
                            f'{MARKET_LABELS[mkt]} / {ch} ({goal}): CPC €{cpc:.2f}, CTR {ctr*100:.1f}%, Click→Session {c2s*100:.0f}%'
                        )
                    else:
                        cpm  = _get('cpm',       mkt, ch, goal) or b.get('cpm', 0)
                        ctr  = _get('ctr',       mkt, ch, goal) or b.get('ctr', 0)
                        freq = _get('frequency', mkt, ch, goal) or b.get('frequency', 3.0)
                        line = f'{MARKET_LABELS[mkt]} / {ch} ({goal}): CPM €{cpm:.2f}, CTR {ctr*100:.2f}%, Freq {freq:.1f}'
                        if ch == 'YouTube':
                            vr = _get('view_rate', mkt, ch, goal) or b.get('view_rate', 0.31)
                            line += f', View Rate {vr*100:.0f}%'
                        bench_lines.append(line)

        channels_in_use = sorted({ch for chs in ai_data['goal_channels'].values() for ch in chs})
        return bench_lines, channels_in_use

    if st.button('Generate Benchmark Explanations', key='btn_bench'):
        api_key = get_api_key()
        if not api_key:
            st.error('No API key found in .streamlit/secrets.toml')
        else:
            bench_lines, channels_in_use = _bench_context()
            prompt = f"""You are a senior paid media strategist explaining benchmark values to a {audience_type} client in the {industry} sector.

The plan uses these channels: {', '.join(channels_in_use) if channels_in_use else 'various digital channels'}.

Benchmarks in use:
{chr(10).join(bench_lines)}

Write a short explanation (150–180 words) covering:
- What CPM / CPC levels mean in the context of {audience_type} {industry} campaigns — why some markets cost more and whether that premium makes sense for this sector
- Why CTR and View Rate vary between markets and channels — relate this to {audience_type} audience behaviour (e.g. professional vs consumer mindset, intent signals)
- What Frequency means for brand recall in {industry}, and how it should be managed differently depending on whether the goal is awareness or conversion

Write in plain English. No jargon. One paragraph per topic. Frame everything for someone who understands the {industry} business, not a media buyer.

After the explanation, add a 'Sources' section listing 2–4 of the most relevant URLs from the reference list below, copied exactly — do not modify or invent any URL.

Reference sources:
{_BENCHMARK_SOURCES}"""
            import anthropic as _anthropic
            client = _anthropic.Anthropic(api_key=api_key)
            with st.spinner('Thinking...'):
                msg = client.messages.create(
                    model='claude-sonnet-4-6',
                    max_tokens=600,
                    messages=[{'role': 'user', 'content': prompt}]
                )
            st.markdown(msg.content[0].text)

    st.divider()
    st.markdown('#### Chat about the benchmarks')
    st.caption('Ask anything — follow up, dig deeper, change direction. The conversation keeps its memory.')

    # Init chat history
    if 'bench_chat' not in st.session_state:
        st.session_state.bench_chat = []

    # Clear button
    if st.session_state.bench_chat:
        if st.button('🗑 Clear conversation', key='bench_clear'):
            st.session_state.bench_chat = []
            st.rerun()

    # Render existing messages
    for msg in st.session_state.bench_chat:
        with st.chat_message(msg['role']):
            st.markdown(msg['content'])

    # Chat input
    if user_input := st.chat_input(
        'e.g. Why is Germany CPM higher than Poland? Can I adjust CTR for a B2B audience?',
        key='bench_chat_input',
    ):
        # Show user message immediately
        with st.chat_message('user'):
            st.markdown(user_input)

        api_key = get_api_key()
        if not api_key:
            with st.chat_message('assistant'):
                st.error('No API key found in .streamlit/secrets.toml')
        else:
            bench_lines, channels_in_use = _bench_context()
            system_ctx = (
                f'You are a senior paid media strategist advising a {audience_type} client '
                f'in the {industry} sector. '
                f'The plan uses these channels: {", ".join(channels_in_use) if channels_in_use else "various digital channels"}. '
                f'Benchmarks in use:\n{chr(10).join(bench_lines)}\n\n'
                f'Answer in 80–150 words. Be direct, reference actual numbers, '
                f'frame for {audience_type} {industry}. No bullet points.'
            )
            # First user turn carries the full context; follow-ups are plain
            api_messages = []
            for i, m in enumerate(st.session_state.bench_chat):
                if i == 0 and m['role'] == 'user':
                    api_messages.append({'role': 'user', 'content': system_ctx + '\n\n---\n\n' + m['content']})
                else:
                    api_messages.append(m)
            # Current turn
            first_turn = len(st.session_state.bench_chat) == 0
            api_messages.append({
                'role': 'user',
                'content': (system_ctx + '\n\n---\n\n' + user_input) if first_turn else user_input,
            })

            import anthropic as _anthropic
            client = _anthropic.Anthropic(api_key=api_key)
            with st.chat_message('assistant'):
                with st.spinner(''):
                    resp = client.messages.create(
                        model='claude-sonnet-4-6',
                        max_tokens=400,
                        messages=api_messages,
                    )
                reply = resp.content[0].text
                st.markdown(reply)

            # Persist both turns
            st.session_state.bench_chat.append({'role': 'user',      'content': user_input})
            st.session_state.bench_chat.append({'role': 'assistant', 'content': reply})
